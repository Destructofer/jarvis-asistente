"""Pruebas de la lógica de Jarvis que no necesitan micrófono, bocinas ni internet.

    .venv\\Scripts\\python -m unittest discover -s tests -v

Cubren los errores que ya pasaron en la vida real (atajos que bloqueaban la PC, "Adiós" que no
cerraba, confirmaciones en cadena, fallos de rutina mal detectados...) para que no regresen.
"""
import datetime
import os
import sys
import tempfile
import threading
import time
import unittest
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import archivos  # noqa: E402
import cerebro  # noqa: E402
import configuracion  # noqa: E402
import control  # noqa: E402
import escuchar  # noqa: E402
import genesis  # noqa: E402
import graph  # noqa: E402
import mantenimiento  # noqa: E402
import presentacion  # noqa: E402
import skills  # noqa: E402
import voz  # noqa: E402


class Registro(unittest.TestCase):
    def test_cada_skill_apunta_a_su_funcion(self):
        """Un @skill pegado a la función equivocada (pasó al insertar una función entre el
        decorador y la suya) solo se notaba al usarla en vivo."""
        import expositor  # noqa: F401
        import graph  # noqa: F401
        import multimedia  # noqa: F401
        import teams  # noqa: F401
        mal = [n for n, s in skills._SKILLS.items()
               if s["fn"].__name__ not in (n, "skill_" + n)]
        self.assertEqual(mal, [])

    def test_parametros_coinciden_con_la_funcion(self):
        import inspect
        for n, s in skills._SKILLS.items():
            firma = inspect.signature(s["fn"]).parameters
            for p in s["params"]:
                self.assertIn(p, firma, f"{n}: el parámetro '{p}' no existe en la función")


class Atajos(unittest.TestCase):
    def test_frases_largas_ya_no_disparan_atajos(self):
        for frase in ["¿Por qué se bloquea el equipo tan seguido?",
                      "Recuérdame a qué hora es la junta con el profe",
                      "Mándale un pantallazo a Juan por WhatsApp",
                      "¿Qué día es el examen de cálculo?",
                      "¿Qué hora es en Tokio?",
                      "¿Cómo está el equipo de México?"]:
            self.assertIsNone(genesis.buscar_atajo(frase), frase)

    def test_frases_exactas_si_disparan(self):
        self.assertEqual(genesis.buscar_atajo("Jarvis, ¿qué hora es?"), "hora_fecha")
        self.assertEqual(genesis.buscar_atajo("Dime qué hora es, por favor"), "hora_fecha")
        self.assertEqual(genesis.buscar_atajo("Toma una captura de pantalla"), "captura_pantalla")
        self.assertEqual(genesis.buscar_atajo("Bloquea la pantalla"), "bloquear_pantalla")

    def test_presentarse_es_instantaneo(self):
        for frase in ["Jarvis, preséntate con el público", "Preséntate, por favor", "Saluda a todos"]:
            self.assertEqual(genesis.atajo_sistema(frase), ("presentarse_al_publico", {}), frase)
        self.assertIsNone(genesis.atajo_sistema("preséntate con el público y luego abre Chrome"))

    def test_salir_con_acento(self):
        for frase in ["Adiós.", "adios", "Hasta luego.", "Jarvis, adiós"]:
            self.assertIn(genesis._limpia_orden(frase), genesis.SALIDAS, frase)


class Presentacion(unittest.TestCase):
    def setUp(self):
        self._orig = (presentacion.en_curso, presentacion._presentacion_activa)
        presentacion.en_curso = lambda: True
        presentacion._presentacion_activa = lambda: None

    def tearDown(self):
        presentacion.en_curso, presentacion._presentacion_activa = self._orig

    def test_atajos_de_diapositivas(self):
        casos = {
            "siguiente": ("presentacion", {"accion": "siguiente"}),
            "Jarvis, siguiente por favor": ("presentacion", {"accion": "siguiente"}),
            "ve a la diapositiva siete": ("presentacion", {"accion": "ir", "numero": 7}),
            "avanza tres": ("presentacion", {"accion": "siguiente", "numero": 3}),
            "regresa": ("presentacion", {"accion": "anterior"}),
            "quita la pantalla negra": ("presentacion", {"accion": "reanudar"}),
        }
        for frase, esperado in casos.items():
            self.assertEqual(genesis.atajo_presentacion(frase), esperado, frase)

    def test_regresar_a_la_presentacion_la_trae_al_frente(self):
        self.assertEqual(genesis.atajo_presentacion("Jarvis, regresa a la presentación"),
                         ("mostrar_presentacion", {}))

    def test_sin_presentacion_decide_el_modelo(self):
        presentacion.en_curso = lambda: False
        self.assertIsNone(genesis.atajo_presentacion("siguiente"))

    def test_abierta_sin_pantalla_completa_responde_clara(self):
        """Antes iba al modelo, que una vez corrió una rutina por un 'siguiente'."""
        presentacion.en_curso = lambda: False
        presentacion._presentacion_activa = lambda: object()
        self.assertEqual(genesis.atajo_presentacion("siguiente"),
                         ("presentacion", {"accion": "siguiente"}))


class Confirmacion(unittest.TestCase):
    def test_si_en_mexicano(self):
        for frase in ["Sí.", "Simón", "Va", "Órale", "Sale", "Claro que sí", "Por supuesto"]:
            self.assertTrue(skills.es_afirmativo(frase), frase)

    def test_ante_la_duda_no(self):
        for frase in ["No.", "Nel", "Sí, no espera", "", "mmm"]:
            self.assertFalse(skills.es_afirmativo(frase), frase)


class BotonesPeligrosos(unittest.TestCase):
    def test_piden_confirmacion(self):
        for boton in ["Eliminar", "Entregar", "Enviar", "Send", "Submit", "Publicar", "Sign out",
                      "Log out", "Logout", "Cerrar sesión", "Transferir", "Aceptar y pagar",
                      "Unirse ahora", "Eliminarlos", "Comprar"]:
            self.assertTrue(control.es_peligroso(boton), boton)

    def test_no_molestan_con_botones_normales(self):
        for boton in ["Descendente", "Llamadas", "Envíos", "Entregas", "Postres", "Compras",
                      "Iniciar sesión", "Guardar", "Siguiente", "Panel de ventas", "Borrador"]:
            self.assertFalse(control.es_peligroso(boton), boton)


class Rutinas(unittest.TestCase):
    def test_fallos_bien_detectados(self):
        self.assertFalse(control.fallo("No había ningún apagado programado."))
        self.assertTrue(control.fallo("Encontré 'Chrome' pero Windows no me dejó traerla al frente."))
        self.assertTrue(control.fallo("'Word' no se cerró todavía; puede estar preguntando"))
        self.assertTrue(control.fallo("No encontré nada parecido a 'Entrar' en Chrome."))
        self.assertTrue(control.fallo(skills.Fallo("cualquier cosa")))
        self.assertFalse(control.fallo("Listo, en Chrome."))

    def test_rutina_se_detiene_en_el_paso_que_falla(self):
        cfg = {"rutinas": {"demo": [{"abrir_web": {"url": "x"}}, {"clic_en": {"texto": "y"}},
                                    {"decir": "no debe llegar"}]}}
        skills.configurar(cfg)
        llamados = []
        control.ejecutor = lambda n, a: (llamados.append(n), skills.Fallo("No encontré 'y'")
                                         if n == "clic_en" else "ok")[1]
        control.hablar = lambda t: llamados.append("decir")
        r = control.rutina("demo")
        self.assertIsInstance(r, skills.Fallo)
        self.assertEqual(llamados, ["abrir_web", "clic_en"])

    def test_continuar_si_falla(self):
        cfg = {"rutinas": {"demo": [{"clic_en": {"texto": "y"}, "continuar_si_falla": True},
                                    {"abrir_web": {"url": "x"}}]}}
        skills.configurar(cfg)
        llamados = []
        control.ejecutor = lambda n, a: (llamados.append(n), skills.Fallo("no"))[1]
        control.hablar = None
        control.rutina("demo")
        self.assertEqual(llamados, ["clic_en", "abrir_web"])


class ProteccionTerceros(unittest.TestCase):
    def test_lo_que_pidio_el_usuario_no_se_pregunta(self):
        self.assertTrue(genesis._pedido_por_usuario({"texto": "Iniciar sesión"},
                                                    "dale clic a iniciar sesión"))
        self.assertTrue(genesis._pedido_por_usuario({"url": "youtube.com"}, "abre youtube"))
        self.assertTrue(genesis._pedido_por_usuario({"texto": "admin@ejemplo.com", "enter": True},
                                                    "escribe admin arroba ejemplo punto com"))

    def test_lo_que_no_pidio_si(self):
        self.assertFalse(genesis._pedido_por_usuario({"url": "sitio-malo.ru/robar"},
                                                     "¿qué dice el último mensaje?"))
        self.assertFalse(genesis._pedido_por_usuario({"nombre": "powerpoint"},
                                                     "lee la ventana"))

    def test_descripcion_legible(self):
        self.assertEqual(genesis._describir_accion("clic_en", {"texto": "Enviar"}), "pulsar «Enviar»")
        self.assertNotIn("{", genesis._describir_accion("abrir_web", {"url": "x.com"}))


class Archivos(unittest.TestCase):
    def test_lista_blanca(self):
        for nombre in ["informe.pdf", "foto.JPG", "clase.pptx", "notas.txt", "cancion.mp3"]:
            self.assertTrue(archivos.se_puede_abrir(nombre), nombre)
        for nombre in ["instalar.py", "atajo.url", "panel.cpl", "consola.msc", "virus.exe",
                       "script.vbe", "sin_extension"]:
            self.assertFalse(archivos.se_puede_abrir(nombre), nombre)


class Transcripcion(unittest.TestCase):
    def test_alucinaciones_de_whisper(self):
        for t in ["Subtítulos realizados por la comunidad de Amara.org", "¡Gracias por ver el video!",
                  "Suscríbete", "", "  "]:
            self.assertTrue(escuchar._es_alucinacion(t), t)
        for t in ["Siguiente diapositiva", "Abre Chrome", "Gracias, Jarvis"]:
            self.assertFalse(escuchar._es_alucinacion(t), t)

    def test_segmentos_dudosos(self):
        self.assertFalse(escuchar._segmento_valido(0.9, -1.2, 1.1, "men la tu chapa"))
        self.assertFalse(escuchar._segmento_valido(0.1, -0.3, 3.0, "sí sí sí sí sí sí"))
        self.assertTrue(escuchar._segmento_valido(0.05, -0.3, 1.2, "abre la presentación"))

    def test_palabra_de_activacion(self):
        p = ["jarvis", "yarvis"]
        self.assertEqual(escuchar.quitar_activacion("Hey Jarvis, siguiente diapositiva", p),
                         "siguiente diapositiva")
        self.assertEqual(genesis._despues_de_palabra("Oye, Jarvis, abre Chrome", p), "abre Chrome")
        self.assertEqual(genesis._despues_de_palabra("bla bla yarvis siguiente", p), "siguiente")
        self.assertEqual(genesis._despues_de_palabra("sin la palabra", p), "")
        # el nombre al final: antes parecía que solo lo habían llamado
        self.assertEqual(genesis._despues_de_palabra("Oye, abre Chrome, Jarvis", p), "abre Chrome")

    def test_nombre_por_parecido(self):
        p = ["jarvis", "yarvis"]
        self.assertTrue(escuchar._buscar_nombre("yaervis abre spotify", p))
        self.assertTrue(escuchar._buscar_nombre("jarbis pon musica", p))
        self.assertIsNone(escuchar._buscar_nombre("travis scott es un rapero", p))
        self.assertIsNone(escuchar._buscar_nombre("abre los archivos del jardin", p))

    def test_umbral_sobre_el_ruido(self):
        escuchar.NIVELES.clear()
        escuchar.NIVELES.extend([0.02, 0.03, 0.04] * 50)  # ruido de un cuarto real
        self.assertGreater(escuchar.umbral_actual(0.004), 0.04)
        escuchar.NIVELES.clear()

    def test_etiqueta_de_accion_nunca_se_dice(self):
        class Loc:
            texto = ""
            def agregar(self, t): self.texto += t
            def cerrar(self): pass
        for pedazos in (["Listo, abrí Spotify. ", "[ACC", "ION: ejecu", "tando]"],
                        ["Todo bien ", "[ACCI"], ["Cuesta [en dólares] 17. [ACCION: buscando]"]):
            turno = genesis.Turno({"voz_activa": True})
            loc = Loc()
            turno._nueva = lambda loc=loc: loc
            for p in pedazos:
                turno.agregar(p)
            turno.cerrar_ronda()
            self.assertNotIn("ACC", loc.texto.upper())
        self.assertIn("[en dólares]", loc.texto)  # los corchetes normales sí se dicen


class Voz(unittest.TestCase):
    def test_partir_frases(self):
        frases = voz.partir_frases("Hola. Soy Jarvis, el asistente del equipo. Hoy les mostramos el sistema.")
        self.assertTrue(all(len(f) >= 20 for f in frases[:-1]))
        self.assertEqual(" ".join(frases).replace("  ", " "),
                         "Hola. Soy Jarvis, el asistente del equipo. Hoy les mostramos el sistema.")

    def test_limpia_simbolos(self):
        self.assertEqual(voz._limpiar("**Hola** mundo `x`"), "Hola mundo x")

    def test_frase_sin_nada_que_decir_no_tumba_a_edge(self):
        antes = dict(voz._fallos)
        clip = voz._Clip("...")
        voz._generar(clip, voz_natural=False)
        self.assertEqual(list(clip.trozos(limite=1)), [])
        self.assertIsNone(clip.error)
        self.assertEqual(voz._fallos, antes)   # ningún motor "falló"

    def test_locucion_en_orden_y_se_puede_cortar(self):
        dichas = []

        class Salida:
            def start(self): pass
            def write(self, d): time.sleep(len(d) / voz.TASA / 20)
            def stop(self): pass
            def abort(self): pass
            def close(self): pass

        def generar(clip, *_a, **_k):
            clip.poner(np.zeros(voz.TASA // 5, np.int16))
            clip.terminar()

        orig_abrir, orig_generar = voz.Locucion._abrir, voz._generar
        voz.Locucion._abrir = lambda self, d: Salida()
        voz._generar = generar
        try:
            loc = voz.Locucion(al_frase=dichas.append)
            for trozo in ["Buenas tardes a todos. ", "Soy Jarvis y hoy ", "les muestro el sistema. ", "Fin."]:
                loc.agregar(trozo)
            loc.cerrar()
            self.assertTrue(loc.esperar(10))
            self.assertEqual(dichas[0], "Buenas tardes a todos.")
            self.assertIn("les muestro el sistema.", " ".join(dichas))
            # corte: lo creado antes de detener() se calla, lo de después suena
            larga = voz.Locucion()
            larga.agregar("Una frase larga que se va a cortar. " * 5)
            voz.detener()
            larga.cerrar()
            larga.esperar(10)
            despues = voz.Locucion(al_frase=dichas.append)
            despues.agregar("Esto sí suena después del corte.")
            despues.cerrar()
            despues.esperar(10)
            self.assertEqual(dichas[-1], "Esto sí suena después del corte.")
        finally:
            voz.Locucion._abrir, voz._generar = orig_abrir, orig_generar


class Respuestas(unittest.TestCase):
    """responder() + Turno con un cerebro falso y una salida de audio falsa."""

    def setUp(self):
        class Salida:
            def start(self): pass
            def write(self, d): pass
            def stop(self): pass
            def abort(self): pass
            def close(self): pass

        def generar(clip, *_a, **_k):
            clip.poner(np.zeros(240, np.int16))
            clip.terminar()

        self.orig = (voz.Locucion._abrir, voz._generar, cerebro.chat, genesis.hud.subtitulo,
                     genesis.hud.estado)
        voz.Locucion._abrir = lambda self, d: Salida()
        voz._generar = generar
        self.dichas = []
        genesis.hud.subtitulo = self.dichas.append
        genesis.hud.estado = lambda *_a: None
        self.cfg = {"voz_activa": True, "respuesta_streaming": True, "skills": {}}
        skills.configurar(self.cfg)

    def tearDown(self):
        (voz.Locucion._abrir, voz._generar, cerebro.chat, genesis.hud.subtitulo,
         genesis.hud.estado) = self.orig

    def _cerebro(self, respuestas):
        llamadas = []

        def chat(cfg, history, tools, temperatura=0.2, al_texto=None, **_k):
            texto = respuestas[len(llamadas)]
            llamadas.append(texto)
            if al_texto:
                for trozo in texto.split(" "):
                    al_texto(trozo + " ")
            return {"content": texto, "tool_calls": [], "origen": "falso"}
        cerebro.chat = chat
        return llamadas

    def test_reintenta_si_llega_en_otro_idioma(self):
        llamadas = self._cerebro(["这是 una respuesta rota.", "Esta es la respuesta buena, en español."])
        turno = genesis.Turno(self.cfg, "¿qué es esto?")
        r = genesis.responder(self.cfg, [{"role": "system", "content": "s"},
                                         {"role": "user", "content": "¿qué es esto?"}], turno=turno)
        turno.cerrar()
        turno.esperar()
        self.assertEqual(len(llamadas), 2)
        self.assertIsInstance(r, genesis.YaDicho)
        self.assertEqual(self.dichas, ["Esta es la respuesta buena, en español."])

    def test_interrumpido_no_vuelve_a_preguntar(self):
        llamadas = self._cerebro(["No debería pedirse."])
        turno = genesis.Turno(self.cfg, "hola")
        turno.interrumpido = True
        with self.assertRaises(genesis.Interrumpido):
            genesis.responder(self.cfg, [{"role": "system", "content": "s"}], turno=turno)
        self.assertEqual(llamadas, [])

    def test_streaming_por_frases(self):
        self._cerebro(["Claro que sí. El sistema tiene tres módulos principales. Y todos funcionan en la nube."])
        turno = genesis.Turno(self.cfg, "explica")
        genesis.responder(self.cfg, [{"role": "system", "content": "s"}], turno=turno)
        turno.cerrar()
        turno.esperar()
        self.assertGreaterEqual(len(self.dichas), 2)
        self.assertEqual(" ".join(self.dichas),
                         "Claro que sí. El sistema tiene tres módulos principales. Y todos funcionan en la nube.")


class Conversacion(Respuestas):
    """Modo conversación: tras contestar sigue escuchando sin "Jarvis"; lo que no era para él
    (le hablaste al público) no se contesta."""

    def _cerebro_trozos(self, trozos):
        def chat(cfg, history, tools, temperatura=0.2, al_texto=None, **_k):
            for t in trozos:
                if al_texto:
                    al_texto(t)
            return {"content": "".join(trozos), "tool_calls": [], "origen": "falso"}
        cerebro.chat = chat

    def _responder(self):
        turno = genesis.Turno(self.cfg, "y luego les muestro", seguimiento=True)
        r = genesis.responder(self.cfg, [{"role": "system", "content": "s"}], turno=turno)
        turno.cerrar()
        turno.esperar()
        return r

    def test_no_era_para_el_se_calla(self):
        self._cerebro_trozos(["<ign", "orar>"])  # la marca puede llegar partida
        self.assertIsInstance(self._responder(), genesis.Ignorado)
        self.assertEqual(self.dichas, [])

    def test_si_era_para_el_contesta(self):
        self._cerebro_trozos(["Cl", "aro, ahí ", "va la respuesta."])
        r = self._responder()
        self.assertNotIsInstance(r, genesis.Ignorado)
        self.assertEqual(self.dichas, ["Claro, ahí va la respuesta."])

    def test_respuesta_muy_corta_no_se_pierde(self):
        self._cerebro_trozos(["Sí."])
        self._responder()
        self.assertEqual(self.dichas, ["Sí."])

    def test_ventana_escucha_sin_palabra(self):
        orig = (escuchar.grabar, escuchar.transcribir, genesis.hud.estado)
        frases = iter(["¿Y cuántos módulos tiene?", "Jarvis, abre Chrome"])
        escuchar.grabar = lambda *a, **k: np.zeros(1600, np.float32)
        escuchar.transcribir = lambda *a, **k: next(frases)
        genesis.hud.estado = lambda *a: None
        cfg = {"palabra_activacion": True, "conversacion_seg": 20,
               "palabras_activacion": ["jarvis"]}
        try:
            genesis.abrir_conversacion(cfg)
            self.assertEqual(genesis.obtener_entrada(cfg), ("¿Y cuántos módulos tiene?", False))
            self.assertTrue(genesis._entrada["seguimiento"])
            genesis.abrir_conversacion(cfg)
            self.assertEqual(genesis.obtener_entrada(cfg), ("abre Chrome", False))
            self.assertFalse(genesis._entrada["seguimiento"])  # dijo su nombre: sí era para él
        finally:
            escuchar.grabar, escuchar.transcribir, genesis.hud.estado = orig
            genesis.cerrar_conversacion()

    def test_clasificar_sin_ia(self):
        para_jarvis = ["¿puedes explicar el de ventas?", "ahora muéstrales cómo se agenda una cita",
                       "oye, ¿y tú cuántas personas ves en la sala?", "abre la presentación",
                       "¿qué opinas de esa pregunta?", "explícales el módulo de reportes",
                       "y ve a inventario"]
        para_publico = ["como pueden ver, compañeros, esto nos ahorró mucho tiempo",
                        "gracias por venir, empecemos con el problema que resolvemos",
                        "les voy a mostrar la siguiente parte del sistema",
                        "nosotros lo construimos en dos días", "bienvenidos a nuestra presentación"]
        dudosas = ["y esto lo construimos en solo dos días durante el hackathon", "¿y qué día es hoy?"]
        for f in para_jarvis:
            self.assertEqual(genesis.clasificar_seguimiento(f), "jarvis", f)
        for f in para_publico:
            self.assertEqual(genesis.clasificar_seguimiento(f), "publico", f)
        for f in dudosas:
            self.assertEqual(genesis.clasificar_seguimiento(f), "duda", f)

    def test_cierre_y_desactivado(self):
        self.assertIn(genesis._limpia_orden("Gracias, Jarvis"), genesis.CIERRE)
        genesis.abrir_conversacion({"palabra_activacion": True, "conversacion_seg": 0})
        self.assertFalse(genesis.en_conversacion())


class Observador(Respuestas):
    """Observa al público: dudas, pausas del expositor, complementos."""

    def test_lee_lo_que_ve(self):
        import observador
        o = observador.parsear('Aquí va: {"escena": "Sala con 30 personas", "personas": "30", '
                               '"platicando_de_frente": "true", "duda": true, "confianza_duda": 0.8, '
                               '"atencion": "Alta"} fin')
        self.assertEqual((o["personas"], o["platicando_de_frente"], o["duda"], o["atencion"]),
                         (30, True, True, "alta"))
        self.assertIsNone(observador.parsear("no es json"))

    def test_cuando_intervenir(self):
        import observador
        conf = {"umbral_duda": 0.7, "pausa_entre_intervenciones_seg": 60, "ventana_seg": 25}
        ahora = 1000.0
        duda = lambda c: {"duda": True, "confianza_duda": c}  # noqa: E731
        nada = {"duda": False, "confianza_duda": 0.0}
        self.assertFalse(observador.debe_intervenir([(ahora, duda(0.75))], ahora, 0, conf))  # gesto suelto
        self.assertTrue(observador.debe_intervenir([(ahora - 8, duda(0.7)), (ahora, duda(0.75))], ahora, 0, conf))
        self.assertTrue(observador.debe_intervenir([(ahora, duda(0.9))], ahora, 0, conf))   # muy claro
        self.assertFalse(observador.debe_intervenir([(ahora, duda(0.9))], ahora, ahora - 30, conf))  # muy seguido
        self.assertFalse(observador.debe_intervenir([(ahora - 8, duda(0.9)), (ahora, nada)], ahora, 0, conf))
        self.assertFalse(observador.debe_intervenir([(ahora - 60, duda(0.95))], ahora, 0, conf))  # ya viejo

    def test_contexto_de_la_camara(self):
        import observador
        observador._historial.clear()
        observador._historial.append((time.time(), {"platicando_de_frente": True}))
        self.assertIn("de frente", observador.contexto_marca())
        observador._historial.clear()
        self.assertEqual(observador.contexto_marca(), "")

    def test_detecta_pausas(self):
        mic = escuchar.MIC
        orig = list(mic._anillo)
        try:
            ahora = time.time()
            mic._anillo.clear()
            for i in range(25):  # 2 s de silencio
                mic._anillo.append((ahora - 2 + i * escuchar.SEG_BLOQUE, np.zeros(escuchar.BLOQUE, np.float32)))
            self.assertTrue(escuchar.en_pausa(1.3, umbral=0.004))
            mic._anillo.append((ahora, np.full(escuchar.BLOQUE, 0.05, np.float32)))  # habló
            self.assertFalse(escuchar.en_pausa(1.3, umbral=0.004))
        finally:
            mic._anillo.clear()
            mic._anillo.extend(orig)

    def test_no_se_graba_a_si_mismo(self):
        import queue as q
        sub = q.Queue()
        for _ in range(30):
            sub.put((time.time(), np.full(escuchar.BLOQUE, 0.05, np.float32)))
        voz.HABLANDO.set()
        try:
            self.assertIsNone(escuchar._grabar_de_sub(sub, 0.004, 0.3, 5, espera_seg=0.5))
        finally:
            voz.HABLANDO.clear()

    def test_complemento_se_guarda_para_la_pausa(self):
        def chat(cfg, history, tools, temperatura=0.2, al_texto=None, **_k):
            for t in ["<compl", "ementar> Si me permites agregar, ", "también funciona sin internet."]:
                al_texto(t)
            return {"content": "", "tool_calls": [], "origen": "falso"}
        cerebro.chat = chat
        turno = genesis.Turno(self.cfg, "como pueden ver", seguimiento=True)
        r = genesis.responder(self.cfg, [{"role": "system", "content": "s"}], turno=turno)
        turno.cerrar()
        turno.esperar()
        self.assertIsInstance(r, genesis.Complemento)
        self.assertEqual(str(r), "Si me permites agregar, también funciona sin internet.")
        self.assertEqual(self.dichas, [])  # nada se dijo todavía

    def test_intervenir_escucha_la_respuesta(self):
        orig = genesis.decir
        genesis.decir = lambda *a, **k: None
        try:
            genesis._notas.clear()
            genesis.intervenir({"palabra_activacion": True, "conversacion_seg": 20}, "¿Te quedó alguna duda?")
            self.assertEqual(genesis._notas[-1]["content"], "¿Te quedó alguna duda?")
            self.assertTrue(genesis.en_conversacion())
            self.assertTrue(escuchar.CONVERSAR.is_set())
        finally:
            genesis.decir = orig
            genesis._notas.clear()
            genesis.cerrar_conversacion()
            escuchar.CONVERSAR.clear()


class Confirmaciones(Respuestas):
    """Lo que pasó en la prueba real: tras leer la pantalla quiso pulsar «EMBIO 2.0», le dijeron
    que no y lo volvió a pedir una y otra vez, tragándose las órdenes siguientes."""

    def test_si_no_o_ninguna(self):
        self.assertTrue(skills.respuesta_si_no("Sí, hazlo"))
        self.assertFalse(skills.respuesta_si_no("No sé..."))
        self.assertIsNone(skills.respuesta_si_no("Dime qué estamos viendo"))
        self.assertIsNone(skills.respuesta_si_no(""))

    def test_un_no_termina_la_orden(self):
        llamadas = []

        def chat(cfg, history, tools, temperatura=0.2, al_texto=None, **_k):
            llamadas.append(1)  # el modelo insiste con lo mismo
            return {"content": "", "tool_calls": [{"id": f"c{len(llamadas)}", "name": "clic_en",
                                                   "args": {"texto": "EMBIO 2.0"}}], "origen": "falso"}
        cerebro.chat = chat
        orig = (genesis._preguntar, genesis.ejecutar_herramienta)
        preguntas = []
        genesis._preguntar = lambda cfg, p: (preguntas.append(p), False)[1]
        genesis.ejecutar_herramienta = lambda cfg, n, a: "Leí la ventana." if n == "leer_ventana" else "ok"
        try:
            # primero lee la pantalla (contenido externo), luego quiere pulsar algo que no pediste
            history = [{"role": "system", "content": "s"},
                       {"role": "tool", "id": "x", "name": "leer_ventana", "content": "..."}]
            turno = genesis.Turno(self.cfg, "dime qué estamos viendo")
            orig_ext = skills.es_externo
            skills.es_externo = lambda n: n in ("leer_ventana", "clic_en")
            try:
                r = genesis.responder(self.cfg, history, turno=turno, texto_usuario="dime qué estamos viendo")
            finally:
                skills.es_externo = orig_ext
            turno.cerrar()
            turno.esperar()
            self.assertLessEqual(len(preguntas), 1)  # a lo mucho UNA pregunta, nunca 8
            self.assertLessEqual(len(llamadas), 2)
            self.assertIn("no lo hago", str(r).lower())
        finally:
            genesis._preguntar, genesis.ejecutar_herramienta = orig

    def test_otra_orden_en_vez_de_contestar(self):
        orig = genesis._preguntar

        def preguntar(cfg, p):
            genesis._orden_pendiente.append("Dime qué estamos viendo")
            return False
        genesis._preguntar = preguntar
        try:
            self.assertFalse(genesis.confirmar({}, "¿Lo hago?"))
            self.assertTrue(genesis._confirmacion["negada"])
            self.assertEqual(genesis.obtener_entrada({"palabra_activacion": True}),
                             ("Dime qué estamos viendo", False))
        finally:
            genesis._preguntar = orig
            genesis._orden_pendiente.clear()


class Criterio(unittest.TestCase):
    """Cuándo Jarvis razona a fondo (opinión) y cuándo contesta rápido (orden)."""

    def test_pide_opinion(self):
        for f in ["¿Qué opinas de usar los lentes?", "¿qué le mejorarías al sistema?",
                  "¿cuál es mejor, React o Vue?", "¿por qué falló?", "¿cómo lo ves?",
                  "¿estás de acuerdo con el jurado?", "¿crees que funcione?"]:
            self.assertTrue(genesis.pide_opinion(f), f)
        for f in ["abre la calculadora", "siguiente", "ve a inventario", "¿qué hora es?"]:
            self.assertFalse(genesis.pide_opinion(f), f)

    def test_espera_sugerida(self):
        self.assertEqual(cerebro.espera_sugerida(Exception("Please try again in 2.3s.")), 2.3)
        self.assertAlmostEqual(cerebro.espera_sugerida(Exception("try again in 850ms")), 0.85)
        self.assertIsNone(cerebro.espera_sugerida(Exception("otro error")))

    def test_presentador_en_la_personalidad(self):
        cfg = {"personality": "Trabajas con {presentador}.", "expositor": {"presentador": "Chris"}}
        self.assertEqual(genesis._personalidad(cfg, False), "Trabajas con Chris.")


class Historial(unittest.TestCase):
    def test_compactar_no_deja_herramientas_huerfanas(self):
        h = [{"role": "system", "content": "s"}]
        for i in range(25):
            h += [{"role": "user", "content": f"u{i}"},
                  {"role": "assistant", "content": "", "tool_calls": [{"id": f"c{i}", "name": "x", "args": {}}]},
                  {"role": "tool", "id": f"c{i}", "name": "x", "content": "r" * 1000}]
        genesis._compactar(h)
        self.assertLessEqual(len(h) - 1, genesis.MAX_MENSAJES)
        self.assertEqual(h[1]["role"], "user")
        self.assertTrue(all(len(m["content"]) <= genesis.MAX_TOOL_VIEJO + 20 for m in h if m["role"] == "tool"))


class Cerebro(unittest.TestCase):
    def test_argumentos_y_host(self):
        self.assertEqual(cerebro._args('{"a": 1}'), {"a": 1})
        self.assertEqual(cerebro._args("no es json"), {})
        self.assertEqual(cerebro._args("[1, 2]"), {})
        self.assertEqual(cerebro.host_de("https://api.groq.com/openai/v1"), "api.groq.com")

    def test_error_de_red(self):
        class APIConnectionError(Exception):
            pass
        self.assertTrue(cerebro.es_error_de_red(APIConnectionError()))
        self.assertTrue(cerebro.es_error_de_red(TimeoutError()))
        self.assertFalse(cerebro.es_error_de_red(ValueError()))


class Graph(unittest.TestCase):
    def test_reparto_respeta_el_tope(self):
        for n in range(1, 7):
            topes = graph._repartir(graph.MAX_TOTAL, n)
            self.assertLessEqual(sum(topes), graph.MAX_TOTAL + 800 * n)
            self.assertEqual(topes, sorted(topes, reverse=True))

    def test_docx_incluye_tablas(self):
        import io

        from docx import Document
        d = Document()
        d.add_paragraph("Actividad 1.4: programa")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Criterio", "Puntos"
        t.cell(1, 0).text, t.cell(1, 1).text = "Funciona", "50"
        buf = io.BytesIO()
        d.save(buf)
        texto = graph._texto_docx(buf.getvalue())
        self.assertIn("Actividad 1.4", texto)
        self.assertIn("Criterio | Puntos", texto)
        self.assertIn("Funciona | 50", texto)


class Mantenimiento(unittest.TestCase):
    def test_solo_borra_temporales_viejos(self):
        with tempfile.TemporaryDirectory() as d:
            base = Path(d)
            viejo, nuevo = base / "viejo.tmp", base / "nuevo.tmp"
            viejo.write_bytes(b"x" * 100)
            nuevo.write_bytes(b"y" * 100)
            hace_una_semana = time.time() - 7 * 86400
            os.utime(viejo, (hace_una_semana, hace_una_semana))
            antes = mantenimiento.TEMP_DIRS
            mantenimiento.TEMP_DIRS = [base]
            try:
                mantenimiento._limpiar_temporales_real()
            finally:
                mantenimiento.TEMP_DIRS = antes
            self.assertFalse(viejo.exists())
            self.assertTrue(nuevo.exists())


class Configuracion(unittest.TestCase):
    def test_guardar_campos_no_pierde_lo_demas(self):
        with tempfile.TemporaryDirectory() as d:
            ruta = Path(d) / "config.json"
            configuracion.guardar({"a": 1, "b": "ñ"}, ruta)
            configuracion.guardar_campos({"b": "otra"}, ruta)
            self.assertEqual(configuracion.cargar(ruta), {"a": 1, "b": "otra"})
            self.assertFalse(ruta.with_suffix(".json.tmp").exists())


class Microfono(unittest.TestCase):
    """El cambio automático a respaldo cuando el cable de la videollamada se queda mudo."""

    class Flujo:
        def __init__(self):
            self.active = True

        def start(self): pass
        def abort(self): self.active = False
        def close(self): pass

    def test_pasa_a_respaldo_y_regresa(self):
        mic = escuchar.Microfono()
        mic.silencio_respaldo = 0.3
        p = escuchar._Fuente("principal", 1)
        r = escuchar._Fuente("respaldo", 2)
        p.stream, r.stream = self.Flujo(), self.Flujo()
        mic._fuentes = {"principal": p, "respaldo": r}
        ruido = np.full((escuchar.BLOQUE, 1), 0.01, np.float32)
        ceros = np.zeros((escuchar.BLOQUE, 1), np.float32)
        mic._llega(p, ruido)
        mic._llega(r, ruido)
        mic._revisar()
        self.assertEqual(mic.activa, "principal")
        time.sleep(0.4)          # la llamada se cae: el cable manda puros ceros
        mic._llega(p, ceros)
        mic._llega(r, ruido)
        mic._revisar()
        self.assertEqual(mic.activa, "respaldo")
        mic._llega(p, ruido)     # vuelve la llamada
        mic._revisar()
        self.assertEqual(mic.activa, "principal")

    def test_suscriptores_reciben_solo_lo_activo(self):
        mic = escuchar.Microfono()
        mic._corriendo = True  # sin abrir dispositivos de verdad
        p = escuchar._Fuente("principal", 1)
        r = escuchar._Fuente("respaldo", 2)
        mic._fuentes = {"principal": p, "respaldo": r}
        q = mic.suscribir()
        mic._llega(p, np.full((escuchar.BLOQUE, 1), 0.5, np.float32))
        mic._llega(r, np.full((escuchar.BLOQUE, 1), 0.9, np.float32))
        _t, bloque = q.get_nowait()
        self.assertAlmostEqual(float(bloque[0]), 0.5)
        self.assertTrue(q.empty())
        self.assertEqual(len(mic.reciente(5)), escuchar.BLOQUE)


class Cognicion(unittest.TestCase):
    """La parte pensante: cuánto pensar, calcular con exactitud y qué aprender del usuario."""

    def test_nivel_de_pensamiento(self):
        import cognicion
        self.assertEqual(cognicion.nivel("abre spotify"), "rapido")
        self.assertEqual(cognicion.nivel("sube el volumen a 60"), "rapido")
        self.assertEqual(cognicion.nivel("¿cómo organizo mi día si tengo examen?"), "profundo")
        self.assertEqual(cognicion.nivel("¿qué me conviene, la laptop o la tablet?"), "profundo")
        self.assertEqual(cognicion.nivel("¿quién ganó el mundial pasado?"), "normal")

    def test_calcular_exacto_y_seguro(self):
        import cognicion
        self.assertTrue(cognicion.calcular("200 - (3*45 + 2*12.50)").endswith("= 40"))
        self.assertTrue(cognicion.calcular("1500 * 16%").endswith("= 240"))
        self.assertIsInstance(cognicion.calcular("10/0"), skills.Fallo)
        self.assertIsInstance(cognicion.calcular('__import__("os").system("dir")'), skills.Fallo)
        self.assertIsInstance(cognicion.calcular("9**999"), skills.Fallo)

    def test_calendario(self):
        import cognicion
        self.assertIn("jueves", cognicion.calendario("dia_de_la_semana", "2026-10-15"))
        self.assertIn("82 días", cognicion.calendario("dias_entre", "2026-10-04", "2026-12-25"))
        self.assertIn("miércoles 14 de octubre", cognicion.calendario("sumar_dias", "2026-10-04", dias=10))

    def test_aprende_solo_lo_util_y_seguro(self):
        import cognicion
        import memoria
        guardados = []
        viejo = (memoria.listar_hechos, memoria.agregar_hecho)
        memoria.listar_hechos = lambda: [{"texto": "Le gusta el rock"}]
        memoria.agregar_hecho = guardados.append
        try:
            nuevos = cognicion._guardar(["Le gusta el rock.", "Su contraseña es hola123",
                                         "Su tarjeta es 4152 3135 0000 1234", "Estudia ingeniería"])
        finally:
            memoria.listar_hechos, memoria.agregar_hecho = viejo
        self.assertEqual(nuevos, ["Estudia ingeniería"])  # sin repetidos ni datos sensibles
        self.assertTrue(cognicion.SOBRE_SI.search(skills._norm("me encanta el café de olla")))
        self.assertFalse(cognicion.SOBRE_SI.search(skills._norm("abre spotify")))


class MemoriaYPreferencias(unittest.TestCase):
    """Lo que Jarvis aprende y sigue usando: preferencias, instrucciones y conversaciones."""

    def setUp(self):
        import memoria
        self._dir = tempfile.TemporaryDirectory()
        self._db = memoria.DB_PATH
        memoria.DB_PATH = Path(self._dir.name) / "prueba.db"

    def tearDown(self):
        import memoria
        memoria.DB_PATH = self._db
        self._dir.cleanup()

    def test_categorias(self):
        import preferencias as pr
        self.assertEqual(pr.categoria_de("el navegador"), "navegador")
        self.assertEqual(pr.categoria_de("mi browser"), "navegador")
        self.assertEqual(pr.categoria_de("mi editor de código"), "codigo")
        self.assertEqual(pr.categoria_de("el buscador"), "buscador")

    def test_preferencia_se_usa_y_se_cambia(self):
        import apps
        import preferencias as pr
        abiertos, viejo = [], (apps.buscar_app, os.startfile)
        apps.buscar_app = lambda n: ("Opera GX", r"C:\Opera\opera.exe", 0.9) if "opera" in n.lower() \
            else ("Google Chrome", r"C:\Chrome\chrome.exe", 0.9)
        os.startfile = lambda ruta, *a: abiertos.append((ruta, a[1] if len(a) > 1 else None))
        try:
            self.assertIn("Opera", pr.fijar_preferencia("navegador", "opera"))
            pr.abrir_url("https://ejemplo.com")
            self.assertEqual(abiertos[-1], (r"C:\Opera\opera.exe", "https://ejemplo.com"))
            self.assertEqual(pr.app_para("el navegador")[0], "Opera GX")
            self.assertIn("cambié", pr.fijar_preferencia("navegador", "chrome"))
            pr.abrir_url("https://ejemplo.com")
            self.assertEqual(abiertos[-1][0], r"C:\Chrome\chrome.exe")
            self.assertEqual(len(pr.todas()), 1)  # se reemplaza, no se acumula
        finally:
            apps.buscar_app, os.startfile = viejo

    def test_buscador_y_servicios(self):
        import preferencias as pr
        self.assertIn("google.com", pr.url_busqueda("pozole"))
        pr.fijar_preferencia("buscador", "DuckDuckGo")
        self.assertIn("duckduckgo.com", pr.url_busqueda("pozole"))
        pr.fijar_preferencia("música", "YouTube")
        self.assertIn("música: YouTube", pr.contexto().replace("musica", "música"))
        self.assertIsInstance(pr.fijar_preferencia("buscador", "Altavista 3000"), skills.Fallo)

    def test_reglas_sin_duplicar_y_olvidar(self):
        import preferencias as pr
        pr.aprender_regla("Siempre que te pida un resumen, guárdalo en Word")
        pr.aprender_regla("siempre que te pida un resumen guárdalo en Word por favor")
        self.assertEqual(len(pr.reglas()), 1)
        self.assertIn("INSTRUCCIONES PERMANENTES", pr.contexto())
        self.assertIsInstance(pr.aprender_regla("mi contraseña del banco es 1234"), skills.Fallo)
        pr.olvidar_preferencia("resumen word")
        self.assertEqual(pr.reglas(), [])

    def test_recordar_conversacion(self):
        import memoria
        memoria.guardar_mensaje("user", "Busca en internet recetas de pozole")
        memoria.guardar_mensaje("assistant", "Buscando recetas de pozole.")
        memoria.guardar_mensaje("user", "¿De qué hablamos del pozole?")
        memoria.guardar_mensaje("assistant", "No recuerdo haber hablado de pozole.")
        r = memoria.recordar_conversacion(tema="pozole", cuando="hace rato")
        self.assertIn("recetas de pozole", r)
        self.assertNotIn("De qué hablamos", r)  # las preguntas de memoria no tapan lo importante


class Avatares(unittest.TestCase):
    """Los GIF del avatar: categoría por nombre, quitar cualquier fondo liso y medir."""

    def test_categoria_por_nombre(self):
        import avatares
        self.assertEqual(avatares.categoria_por_nombre("celebrando_baile"), "celebrando")
        self.assertEqual(avatares.categoria_por_nombre("ejecutando_3"), "ejecutando")
        self.assertEqual(avatares.categoria_por_nombre("libre_caminando"), "libre")
        self.assertEqual(avatares.categoria_por_nombre("Espera"), "espera")
        self.assertIsNone(avatares.categoria_por_nombre("cartas"))

    def _personaje(self, fondo):
        from PIL import Image, ImageDraw
        im = Image.new("RGB", (120, 160), fondo)
        d = ImageDraw.Draw(im)
        d.ellipse((45, 10, 75, 40), fill=(240, 200, 150))
        d.rectangle((48, 42, 72, 110), fill=(200, 40, 40))
        d.rectangle((20, 120, 40, 150), fill=(90, 90, 90))  # un objeto más chico al lado
        return im

    def test_quita_fondos_de_color(self):
        import avatares
        for fondo in ((170, 210, 245), (255, 255, 255), (51, 51, 51)):
            a = avatares.sin_fondo(self._personaje(fondo)).getchannel("A")
            self.assertEqual(a.getpixel((2, 2)), 0, fondo)       # el fondo, transparente
            self.assertEqual(a.getpixel((60, 80)), 255, fondo)   # el cuerpo, entero

    def test_respeta_gif_transparente(self):
        from PIL import Image
        import avatares
        im = Image.new("RGBA", (50, 50), (0, 0, 0, 0))
        im.paste((10, 200, 10, 255), (15, 15, 35, 35))
        a = avatares.sin_fondo(im).getchannel("A")
        self.assertEqual((a.getpixel((2, 2)), a.getpixel((25, 25))), (0, 255))

    def test_mide_al_personaje_no_al_objeto(self):
        import avatares
        alto, pies, centro = avatares._medir(avatares.sin_fondo(self._personaje((170, 210, 245))))
        self.assertGreater(alto, 90)          # cabeza a cuerpo (unidos), no el objeto de 30 px
        self.assertAlmostEqual(centro, 60, delta=6)


class LectorDeCamara(unittest.TestCase):
    """La cámara que se cae, se bloquea o la toma otra app vuelve sola (antes gestos, presencia
    y la realidad aumentada se quedaban "trabados" hasta reiniciar Jarvis)."""

    def setUp(self):
        import cv2
        import camara
        self.cv2, self.camara = cv2, camara
        self._orig = cv2.VideoCapture
        self.aperturas = []

    def tearDown(self):
        self.cv2.VideoCapture = self._orig

    def _camara_falsa(self, comportamiento):
        prueba = self

        class Falsa:
            def __init__(self, *a):
                self.n = len(prueba.aperturas)
                prueba.aperturas.append(self)
                self.lecturas = 0

            def isOpened(self):
                return True

            def set(self, *a):
                return True

            def read(self):
                self.lecturas += 1
                return comportamiento(self)

            def release(self):
                pass
        self.cv2.VideoCapture = Falsa

    def _esperar_imagen_nueva(self, lector, desde, limite=12.0):
        fin = time.time() + limite
        while time.time() < fin:
            cuadro, ts = lector.ultimo(0)
            if cuadro is not None and ts > desde:
                return True
            time.sleep(0.05)
        return False

    def test_se_reabre_si_deja_de_dar_imagen(self):
        cuadro = np.zeros((4, 4, 3), np.uint8)

        def comp(cam):  # la primera cámara da 5 cuadros y luego nada; la segunda funciona
            time.sleep(0.02)
            return (True, cuadro) if (cam.n > 0 or cam.lecturas <= 5) else (False, None)
        self._camara_falsa(comp)
        lector = self.camara._Lector()
        self.assertTrue(self._esperar_imagen_nueva(lector, 0))
        t = time.time()
        self.assertTrue(self._esperar_imagen_nueva(lector, t + 0.5))
        self.assertGreaterEqual(len(self.aperturas), 2)
        lector.parar()

    def test_lectura_bloqueada_se_abandona(self):
        cuadro = np.zeros((4, 4, 3), np.uint8)
        bloqueo = threading.Event()

        def comp(cam):  # la primera se cuelga en la lectura (como cuando otra app la toma)
            if cam.n == 0 and cam.lecturas > 3:
                bloqueo.wait(30)
                return False, None
            time.sleep(0.02)
            return True, cuadro
        self._camara_falsa(comp)
        lector = self.camara._Lector()
        self.assertTrue(self._esperar_imagen_nueva(lector, 0))
        t = time.time()
        self.assertTrue(self._esperar_imagen_nueva(lector, t + 0.5))
        self.assertGreaterEqual(len(self.aperturas), 2)
        bloqueo.set()
        lector.parar()


class InteraccionConPantalla(unittest.TestCase):
    """Elegir lo que se ve ("el primer video", "la segunda playlist") y controlar el video."""

    def setUp(self):
        import interaccion
        self.I = interaccion

    def _cosa(self, nombre, url, x, y, visible=True, host="www.youtube.com"):
        c = self.I.Cosa(None, nombre, "Hyperlink", (x, y, x + 300, y + 170), url)
        c.clase = self.I.clasificar(url, nombre, "Hyperlink", host)
        c.visible = visible
        return c

    def test_clasifica_por_la_url(self):
        c = self.I.clasificar
        yt = "https://www.youtube.com"
        self.assertEqual(c(yt + "/watch?v=abcdefghijk"), "video")
        # un video ofrecido "como radio" sigue siendo ese video
        self.assertEqual(c(yt + "/watch?v=abcdefghijk&list=RDabcdefghijk&start_radio=1", "Canción"), "video")
        self.assertEqual(c(yt + "/watch?v=C86dq7o5Hpo&list=RDGMEMR48zJN", "Mix: Trap"), "mix")
        self.assertEqual(c(yt + "/watch?v=abcdefghijk&list=PLx0sYbCqOb8"), "playlist")
        self.assertEqual(c(yt + "/playlist?list=PLx0sYbCqOb8"), "playlist")
        self.assertEqual(c(yt + "/shorts/abcdefghijk"), "short")
        self.assertEqual(c(yt + "/@SadLyrics-Traducciones"), "canal")
        self.assertEqual(c(yt + "/feed/subscriptions"), "navegacion")
        self.assertEqual(c("https://open.spotify.com/track/123"), "cancion")
        self.assertEqual(c("https://open.spotify.com/album/123"), "album")
        g = "www.google.com"
        self.assertEqual(c("https://es.wikipedia.org/wiki/Python", "Python", "Hyperlink", g), "resultado")
        self.assertEqual(c("https://www.google.com/search?q=x&tbm=isch", "Imágenes", "Hyperlink", g), "navegacion")
        self.assertEqual(c("", "Pausa (k)", "Button"), "boton")

    def test_pedidos_con_varios_tipos(self):
        self.assertEqual(self.I.clases_pedidas("canción o playlist"),
                         {"cancion", "video", "playlist", "mix", "album"})
        self.assertEqual(self.I.clases_pedidas("videos"), {"video"})
        self.assertIn("video", self.I.clases_pedidas(""))   # sin tipo: lo primero que se vea

    def test_elige_en_orden_de_lectura(self):
        v = "https://www.youtube.com/watch?v="
        cosas = self.I.ordenar([
            self._cosa("Segundo renglón", v + "bbbbbbbbbbb", 20, 400),
            self._cosa("Mix: Pop", v + "ccccccccccc&list=RDGMEMQ1dJ7w", 700, 120),
            self._cosa("Primero 3 minutos y 7 segundos", v + "aaaaaaaaaaa", 20, 125),
            self._cosa("Abajo, sin verse", v + "ddddddddddd", 20, 1900, visible=False),
        ])
        e = self.I.escoger
        clases = self.I.clases_pedidas
        self.assertEqual(e(cosas, clases("video"), 1).nombre, "Primero 3 minutos y 7 segundos")
        self.assertEqual(e(cosas, clases("cancion o playlist"), 2).nombre, "Mix: Pop")
        self.assertEqual(e(cosas, clases("playlist"), 1).nombre, "Mix: Pop")
        self.assertEqual(e(cosas, clases("video"), -1).nombre, "Segundo renglón")  # el último que se ve
        self.assertEqual(e(cosas, clases("video"), 3).nombre, "Abajo, sin verse")   # más de los que se ven
        self.assertIsNone(e(cosas, clases("short"), 1))
        self.assertEqual(e(cosas, clases("video"), 1, contiene="segundo").nombre, "Segundo renglón")
        self.assertEqual(self.I.limpio("Primero 3 minutos y 7 segundos"), "Primero")

    def test_miniatura_y_titulo_son_uno(self):
        url = "https://www.youtube.com/watch?v=aaaaaaaaaaa&pp=xyz"
        unidos = self.I._unir([self._cosa("", url, 20, 120),
                               self._cosa("Título completo", url.replace("xyz", "abc"), 20, 300)])
        self.assertEqual(len(unidos), 1)
        self.assertEqual(unidos[0].nombre, "Título completo")
        self.assertEqual(unidos[0].rect[1], 120)   # en el lugar de la miniatura

    def test_boton_omitir_no_confunde_omitir_navegacion(self):
        o = self.I.OMITIR
        for nombre in ("Omitir", "Omitir anuncio", "Skip Ad", "Saltar anuncios", "Omitir ›"):
            self.assertTrue(o.match(nombre), nombre)
        for nombre in ("Omitir navegación", "Omitir anuncio en 5", "Anuncio"):
            self.assertFalse(o.match(nombre), nombre)

    def test_atajos_de_reproduccion(self):
        a = genesis.atajo_medios
        self.assertEqual(a("pausa"), ("controlar_reproduccion", {"accion": "pausar"}))
        self.assertEqual(a("Jarvis, salta el anuncio por favor"),
                         ("controlar_reproduccion", {"accion": "saltar_anuncio"}))
        self.assertEqual(a("siguiente canción")[1]["accion"], "siguiente")
        self.assertEqual(a("adelanta 30 segundos")[1], {"accion": "adelantar", "segundos": 30})
        self.assertEqual(a("regresa un minuto")[1], {"accion": "retroceder", "segundos": 60})
        self.assertIsNone(a("regresa"))       # eso es "atrás" en el navegador
        self.assertIsNone(a("abre youtube"))
        self.assertEqual(genesis.buscar_atajo("¿qué canción es esta?"), "que_suena")

    def test_sin_llamarlo_no_contesta_platica_ajena(self):
        cfg = genesis.load_config() if (Path(genesis.__file__).parent / "config.json").exists() else {}
        orig = (genesis.cerebro.nube_disponible, genesis.escuchar.frase_de_la_pc)
        try:
            genesis.escuchar.frase_de_la_pc = lambda minimo=0.5: False
            genesis.cerebro.nube_disponible = lambda c: True
            self.assertTrue(genesis._no_es_para_mi(cfg, "Sí, sí."))             # corta, no pide nada
            self.assertEqual(genesis._no_es_para_mi(cfg, "pausa"), "")           # atajo: sí es para él
            self.assertEqual(genesis._no_es_para_mi(cfg, "muchas gracias"), "")  # cierra la plática
            largo = "uno es más sintético que el otro, dejas de hacer cosas"
            self.assertEqual(genesis._no_es_para_mi(cfg, largo), "")             # con nube decide el modelo
            genesis.cerebro.nube_disponible = lambda c: False
            self.assertTrue(genesis._no_es_para_mi(cfg, largo))                  # sin nube: callado
            genesis.cerebro.nube_disponible = lambda c: True
            genesis.escuchar.frase_de_la_pc = lambda minimo=0.5: True
            self.assertTrue(genesis._no_es_para_mi(cfg, largo))                  # sonaba la PC
        finally:
            genesis.cerebro.nube_disponible, genesis.escuchar.frase_de_la_pc = orig

    def test_abrir_y_reproducir_es_de_varios_pasos(self):
        t = skills._norm("abre youtube y reproduce la primera cancion o playlist que veas")
        self.assertTrue(genesis.VARIOS_PASOS.search(t))
        nombres = genesis.elegir_herramientas(t, [])
        self.assertTrue({"youtube", "elegir_en_pantalla", "controlar_reproduccion"} <= nombres)

    def test_elige_la_sesion_que_corresponde(self):
        sonando = {"app": "Spotify.exe", "estado": 4, "actual": False}
        pausada = {"app": "MSEdge", "estado": 5, "actual": True}
        e = self.I._elegir_sesion
        self.assertIs(e([pausada, sonando], "pausar"), sonando)    # pausa lo que suena
        self.assertIs(e([pausada, sonando], "reanudar"), pausada)  # reanuda lo pausado
        self.assertIsNone(e([], "pausar"))


class RealidadAumentada(unittest.TestCase):
    """La física y los gestos del modo realidad aumentada, sin cámara ni ventanas reales."""

    def setUp(self):
        import realidad
        self.R = realidad
        self.ruedas, self.clics, self.escrito, self.teclas = [], [], [], []
        # Nada de entrada ni ventanas reales: todo lo que tocaría Windows queda anotado
        self._nombres = ("_cursor", "_boton_izq", "_rueda", "_boton_der", "_subir", "_escribir",
                         "_tecla", "_dar_foco")
        self._orig = {n: getattr(realidad, n) for n in self._nombres}
        realidad._cursor = lambda x, y: self.clics.append(("mover", round(x), round(y)))
        realidad._boton_izq = lambda abajo: self.clics.append(("abajo" if abajo else "arriba",))
        realidad._rueda = self.ruedas.append
        realidad._boton_der = lambda: self.clics.append(("derecho",))
        realidad._subir = lambda h: None
        realidad._escribir = self.escrito.append
        realidad._tecla = lambda *t: self.teclas.append(t)
        realidad._dar_foco = lambda h: True
        self.esc = realidad.Escena({})
        self.esc._armar()
        self.esc._recapturar = lambda cambios: None
        self.esc._dimensionar_reales = lambda celdas, area: None
        self.esc._revisar_campo_de_texto = lambda: None

    def tearDown(self):
        for n, f in self._orig.items():
            setattr(self.R, n, f)

    def _ventana(self):
        e = self.R.Elemento("ventana", "Prueba", 640, 360, 300, 210, self.R._icono(None, 22, "P"), hwnd=1)
        self.esc.elementos.insert(0, e)
        return e

    def _mano(self, x, y):
        m = self.esc.manos.setdefault("Right", self.R.Mano())
        m.x, m.y, m.visto = x, y, time.time() + 999
        return m

    def _correr(self, n, dt=1 / 30):
        for _ in range(n):
            self.esc._animar(dt, time.time())

    def test_resorte_llega_al_objetivo(self):
        x, v = 0.0, 0.0
        maximo = 0.0
        for _ in range(60):
            x, v = self.R._resorte(x, v, 100.0, 1 / 30, 230, 21)
            maximo = max(maximo, x)
        self.assertAlmostEqual(x, 100.0, delta=0.5)
        self.assertLess(maximo, 115)  # rebota un poco, no se dispara
        # un cuadro muy lento no lo vuelve inestable
        x, v = self.R._resorte(0.0, 0.0, 100.0, 0.5, 230, 21)
        self.assertLess(abs(x), 200)

    def test_quieto_no_se_deforma(self):
        e = self._ventana()
        self._correr(40)
        self.assertIsNone(e.deformacion(*e.rect(), time.time()))

    def test_arrastrar_a_la_derecha_inclina_hacia_alla(self):
        e = self._ventana()
        self._correr(30)
        x, y, w, h = e.rect()
        m = self._mano(x + w / 2, y + 10)
        self.esc._presionar("Right", m.x, m.y)
        for _ in range(8):
            m.x += 25
            self.esc._mover("Right", m.x, m.y)
            self._correr(1)
        self.assertGreater(e.vx, 300)
        q = e.deformacion(*e.rect(), time.time())
        self.assertIsNotNone(q)
        izq, der = q[3][1] - q[0][1], q[2][1] - q[1][1]
        self.assertLess(der, izq)       # el borde de adelante (derecho) se va hacia atrás
        self.assertGreater(e.giro, 0)   # y cuelga de la mano como péndulo
        self.assertGreater(e.alzado, 0.5)

    def test_soltar_en_movimiento_la_lanza_sin_sacarla_de_la_pantalla(self):
        e = self._ventana()
        self._correr(30)
        x, y, w, h = e.rect()
        m = self._mano(x + w / 2, y + 10)
        self.esc._presionar("Right", m.x, m.y)
        for _ in range(10):
            m.x += 40
            self.esc._mover("Right", m.x, m.y)
            self._correr(1)
        antes = e.cx
        self.esc._soltar("Right", m.x, m.y)
        self.assertGreater(e.tcx, antes)       # sigue hacia donde iba
        self.assertLessEqual(e.tcx, self.R.W)  # pero no se sale
        self._correr(60)
        self.assertAlmostEqual(e.cx, e.tcx, delta=2)

    def test_icono_del_dock_es_elastico(self):
        app = next(a for a in self.esc.elementos if a.tipo == "app")
        self._correr(60)
        m = self._mano(app.cx, app.cy)
        self.esc._presionar("Right", m.x, m.y)
        m.y -= 200
        self.esc._mover("Right", m.x, m.y)
        self.assertAlmostEqual(app.tcy, app.casa[1] - 80, delta=1)  # se estira menos que la mano
        self.esc._soltar("Right", m.x, m.y)
        self.assertEqual((app.tcx, app.tcy), app.casa)

    def test_boton_de_la_barra_no_arrastra_la_ventana(self):
        e = self._ventana()
        self._correr(30)
        e.botones = {"Reducir": (600, 260, 60, 20)}
        m = self._mano(620, 270)
        self.esc._presionar("Right", 620, 270)
        antes = (e.tcx, e.tcy)
        self.esc._mover("Right", 700, 330)
        self.assertEqual((e.tcx, e.tcy), antes)

    def test_presionar_una_ventana_que_acaba_de_desaparecer(self):
        e = self._ventana()
        self._correr(30)
        bajo = self.esc._bajo
        self.esc._bajo = lambda x, y: e
        self.esc.elementos.remove(e)   # el hilo de ventanas la quitó justo antes
        self._mano(e.cx, e.cy)
        self.esc._presionar("Right", e.cx, e.cy)  # antes: ValueError que cortaba el cuadro
        self.esc._bajo = bajo

    def test_scroll_rapido_sigue_solo_y_se_frena(self):
        e = self._ventana()
        e.interactiva, e.zona = True, (400, 200, 480, 300)
        m = self._mano(640, 350)
        self.esc._presionar("Right", 640, 350)
        for _ in range(5):
            m.y -= 25
            self.esc._mover("Right", m.x, m.y)
        m.vy = -800
        self.esc._soltar("Right", m.x, m.y)
        self.assertIsNotNone(self.esc.inercia)
        n = len(self.ruedas)
        for _ in range(90):
            self.esc._seguir_inercia(1 / 30)
        self.assertGreater(len(self.ruedas), n)
        self.assertTrue(all(u < 0 for u in self.ruedas[n:]))   # misma dirección que la mano
        self.assertIsNone(self.esc.inercia)                    # y se detuvo sola

    def _usable(self, x=640, y=360):
        """Una ventana grande y usable con su imagen ya dibujada (zona)."""
        e = self._ventana()
        e.cx = e.tcx = x
        e.cy = e.tcy = y
        e.escala = e.tescala = 1.6
        e.zona = (x - 230, y - 120, 460, 260)
        self.esc._a_pantalla = lambda v, px, py: (px * 2, py * 2)  # panel -> "pantalla real"
        return e

    def _toque(self, x, y):
        m = self._mano(x, y)
        self.esc._presionar("Right", x, y)
        self.esc._soltar("Right", x, y)
        return m

    def test_tocar_otra_ventana_grande_la_usa_sin_moverla(self):
        a = self._usable(400, 300)
        a.interactiva = True
        b = self._usable(900, 300)
        lugar = (b.tcx, b.tcy, b.tescala)
        self._toque(900, 320)
        self.assertTrue(b.interactiva and not a.interactiva)
        self.assertEqual((b.tcx, b.tcy, b.tescala), lugar)          # no se reacomodó
        self.assertIn(("abajo",), self.clics)                       # y le dio clic ahí mismo

    def test_doble_toque_cae_en_el_mismo_pixel(self):
        e = self._usable()
        e.interactiva = True
        self._toque(600, 350)
        self._toque(603, 352)   # la mano tiembla un poco
        movs = [c for c in self.clics if c[0] == "mover"]
        self.assertEqual(movs[-1], movs[0])    # Windows lo cuenta como doble clic

    def test_arrastrar_hacia_otra_ventana(self):
        a = self._usable(400, 300)
        a.interactiva = True
        b = self._usable(900, 300)
        m = self._mano(420, 320)
        self.esc._presionar("Right", 420, 320)
        m.t0 -= 1                              # mantuvo quieto: agarra
        self.esc._mover("Right", 420, 320)
        self.assertTrue(m.raton_abajo)
        self.esc._mover("Right", 880, 330)     # la mano pasa a la otra ventana
        self.assertIs(m.destino, b)
        self.assertEqual(self.clics[-1], ("mover", 1760, 660))
        self.esc._soltar("Right", 880, 330)
        self.assertEqual(self.clics[-1], ("arriba",))   # soltó ahí

    def test_teclado_escribe_con_acentos_y_mayusculas(self):
        e = self._usable()
        e.interactiva = True
        self.esc._mostrar_teclado(True)
        t = self.esc.teclado
        for clave in ("MAYUS", "h", "ACENTO", "o", "ñ", "ESPACIO", "BORRAR", "ENTER", "COPIAR"):
            self.esc._pulsar_tecla(t, clave)
        self.assertEqual("".join(self.escrito), "Hóñ ")
        self.assertEqual(self.teclas, [("BORRAR",), ("ENTER",), ("CTRL", "C")])
        self.assertFalse(t.mayus or t.acento)   # se usan una vez
        self.esc._pulsar_tecla(t, "OCULTAR")
        self.assertFalse(self.esc._teclado_visible())

    def test_teclas_dentro_del_teclado_y_sin_encimarse(self):
        w, h = 860, 280
        teclas = self.R._disposicion_teclado(w, h)
        claves = [c for _r, c, _e in teclas]
        for k in ("ñ", "ENTER", "BORRAR", "ESPACIO", "MAYUS", "ACENTO", "OCULTAR", "@"):
            self.assertIn(k, claves)
        rects = [r for r, _c, _e in teclas]
        for i, (x, y, tw, th) in enumerate(rects):
            self.assertTrue(0 <= x and x + tw <= w and 0 <= y and y + th <= h)
            for (x2, y2, w2, h2) in rects[i + 1:]:
                self.assertFalse(x < x2 + w2 and x2 < x + tw and y < y2 + h2 and y2 < y + th)

    def test_mosaico_no_encima_ventanas(self):
        for i in range(5):
            e = self._ventana()
            e.t0 = i
            e.hwnd = 10 + i
        self.esc._alternar_mosaico()
        for _ in range(60):
            for e in self.esc._ventanas():
                e.animar(1 / 30)
        grandes = [e for e in self.esc._ventanas() if e.tescala > 0.9]
        self.assertEqual(len(grandes), 4)
        cajas = [e.rect() for e in grandes]
        for i, (x, y, w, h) in enumerate(cajas):
            for (x2, y2, w2, h2) in cajas[i + 1:]:
                self.assertFalse(x < x2 + w2 and x2 < x + w and y < y2 + h2 and y2 < y + h)

    def test_ventanas_vuelven_a_como_estaban(self):
        import json
        guardado = self.R.COLOCACIONES
        self.R.COLOCACIONES = Path(tempfile.mkdtemp()) / "ventanas.json"
        try:
            self.R._guardar_colocaciones({123456789: [[0, 1, [0, 0], [0, 0], [0, 0, 10, 10]], "x.exe"]})
            self.assertTrue(json.loads(self.R.COLOCACIONES.read_text()))
            self.R.restaurar_pendientes()   # la ventana ya no existe: no pasa nada y se limpia
            self.assertFalse(self.R.COLOCACIONES.exists())
        finally:
            self.R.COLOCACIONES = guardado

    # --- salir del modo solo a propósito ---
    def test_toque_rapido_en_escritorio_no_saca_del_modo(self):
        e = self._usable()
        e.interactiva = True
        e.botones = {"Escritorio": (700, 200, 90, 24)}
        self._mano(740, 212)
        self.esc._presionar("Right", 740, 212)
        self.esc._soltar("Right", 740, 212)
        self.assertFalse(self.esc.parar.is_set())
        self.assertIn("Mantén", self.esc.aviso)

    def test_escritorio_sostenido_si_saca_del_modo(self):
        e = self._usable()
        e.interactiva = True
        e.botones = {"Escritorio": (700, 200, 90, 24)}
        m = self._mano(740, 212)
        self.esc._presionar("Right", 740, 212)
        m.t0 -= 1.0
        self.esc._soltar("Right", 740, 212)
        self.assertTrue(self.esc.parar.is_set())
        self.assertIn("botón Escritorio", self.esc.motivo)

    def test_ventana_soltada_de_pasada_en_la_zona_no_saca_del_modo(self):
        e = self._ventana()
        self._correr(20)
        x, y, w, h = e.rect()
        m = self._mano(x + w / 2, y + 10)
        self.esc._presionar("Right", m.x, m.y)
        zx, zy, zw, zh = self.esc._zona_escritorio()
        self.esc._mover("Right", zx + zw / 2, zy + zh / 2)
        self.esc._soltar("Right", zx + zw / 2, zy + zh / 2)   # sin detenerse ahí
        self.assertFalse(self.esc.parar.is_set())
        # sosteniéndola un momento sobre la zona, sí
        self._correr(30)
        x, y, w, h = e.rect()
        self.esc._presionar("Right", x + w / 2, y + 10)
        self.assertIs(m.elem, e)
        self.esc._mover("Right", zx + zw / 2, zy + zh / 2)
        m.t_zona -= 1.0
        self.esc._soltar("Right", zx + zw / 2, zy + zh / 2)
        self.assertTrue(self.esc.parar.is_set())

    # --- cámara con problemas: que se note, no que parezca trabado ---
    def _salud(self, brillo=128.0, sin_imagen=0.0):
        ahora = time.time()
        return {"t": ahora, "n": 30, "brillo": brillo, "nuevo": ahora - sin_imagen, "reabrir": 0.0}

    def test_avisa_si_la_camara_esta_negra(self):
        self.esc._revisar_camara(self._salud(brillo=3), True, time.time(), None)
        self.assertIn("negra", self.esc.aviso)

    def test_reabre_la_camara_si_deja_de_mandar_imagen(self):
        paradas = []

        class Lector:
            def parar(self):
                paradas.append(1)
        self.esc._revisar_camara(self._salud(sin_imagen=3.0), False, time.time(), Lector())
        time.sleep(0.1)
        self.assertIn("dejó de mandar", self.esc.aviso)
        self.assertEqual(paradas, [1])

    def test_avisa_si_la_camara_va_lenta(self):
        salud = self._salud()
        salud["t"], salud["n"] = time.time() - 2, 2   # 1 cuadro por segundo
        self.esc._revisar_camara(salud, True, time.time(), None)
        self.assertIn("lenta", self.esc.aviso)

    def test_con_buena_camara_no_avisa(self):
        self.esc._revisar_camara(self._salud(), True, time.time(), None)
        self.assertEqual(self.esc.aviso, "")

    def test_poca_luz_se_aclara(self):
        oscura = np.full((36, 64, 3), 30, np.uint8)
        self.assertGreater(int(self.R._aclarar(oscura, 30).mean()), 80)
        normal = np.full((36, 64, 3), 120, np.uint8)
        self.assertIs(self.R._aclarar(normal, 120), normal)   # con luz, ni se toca

    # --- clic derecho: solo a propósito ---
    def _resultado(self, pulgar_indice, pulgar_medio):
        """Un resultado de MediaPipe con una mano: distancias relativas al tamaño de la mano."""
        from types import SimpleNamespace as N
        pts = [[0.5, 0.5] for _ in range(21)]
        pts[0], pts[9] = [0.5, 0.7], [0.5, 0.5]          # tamaño de la mano: 0.2
        pts[4] = [0.45, 0.45]
        pts[8] = [0.45 + 0.2 * pulgar_indice, 0.45]
        pts[12] = [0.45, 0.45 + 0.2 * pulgar_medio]
        return N(hand_landmarks=[[N(x=x, y=y) for x, y in pts]],
                 handedness=[[N(category_name="Right")]])

    def _manos(self, pi, pm, segundos, paso=1 / 30):
        t = time.time()
        fin = t + segundos
        while t < fin:
            self.esc._procesar_manos(self._resultado(pi, pm), t)
            t += paso
        return t

    def test_soltar_un_pellizco_no_da_clic_derecho(self):
        derechos = []
        self.esc._clic_derecho = lambda x, y: derechos.append((x, y))
        self._manos(0.1, 0.2, 0.3)               # pellizco normal (índice) con el medio cerca
        self.esc.manos["Right"].t_suelta = time.time()
        self._manos(0.8, 0.2, 0.3)               # suelta: el índice se abre, el medio sigue cerca
        self._manos(0.8, 0.9, 0.2)
        self.assertEqual(derechos, [])

    def test_pulgar_con_medio_sostenido_da_clic_derecho(self):
        derechos = []
        self.esc._clic_derecho = lambda x, y: derechos.append((x, y))
        t = self._manos(0.9, 0.9, 0.2)
        self.esc.manos["Right"].t_suelta = 0.0
        hasta = t + 0.3
        while t < hasta:                          # pulgar con medio, índice abierto, 0.3 s
            self.esc._procesar_manos(self._resultado(0.9, 0.1), t)
            t += 1 / 30
        self.esc._procesar_manos(self._resultado(0.9, 0.9), t)   # y lo suelta
        self.assertEqual(len(derechos), 1)

    def test_captura_cerrada_no_revive(self):
        cap = self.R.Captura(0)
        cap.cerrar()
        cap.iniciar(rapida=True)   # p. ej. un reinicio que llegó tarde desde otro hilo
        self.assertIsNone(cap._ctl)

    def test_dibuja_un_cuadro_con_todo_animandose(self):
        e = self._ventana()
        e.captura = None
        e.mini_pw = np.full((300, 520, 3), 90, np.uint8)
        m = self._mano(e.cx, e.cy - 90)
        self.esc._presionar("Right", m.x, m.y)
        c = np.zeros((self.R.H, self.R.W, 3), np.uint8)
        for _ in range(5):
            m.x += 30
            self.esc._mover("Right", m.x, m.y)
            self.esc._dibujar(c, time.time(), 1 / 30)
        self.assertGreater(int(c.max()), 0)


if __name__ == "__main__":
    unittest.main()


class CodigoLimpio(unittest.TestCase):
    def test_sin_caracteres_de_control_escondidos(self):
        """Un \\b mal escapado quedó como carácter de retroceso: la regex nunca coincidía y nadie
        lo notaba (pasó dos veces). Ningún .py del proyecto debe tener caracteres de control."""
        raiz = Path(__file__).resolve().parents[1]
        malos = []
        for archivo in list(raiz.glob("*.py")) + list((raiz / "tests").glob("*.py")):
            texto = archivo.read_text(encoding="utf-8")
            for n, linea in enumerate(texto.splitlines(), 1):
                if any(ord(c) < 32 and c != "\t" for c in linea):
                    malos.append(f"{archivo.name}:{n}")
        self.assertEqual(malos, [])


class Personalidades(unittest.TestCase):
    def setUp(self):
        import memoria
        import personalidades
        self.P = personalidades
        self._db = memoria.DB_PATH
        memoria.DB_PATH = Path(tempfile.mkdtemp()) / "prueba.db"
        personalidades._estado["cache"] = None
        self.conf = {"edge_voz": "es-MX-JorgeNeural", "edge_velocidad": "+5%", "edge_tono": "+0Hz"}
        self._voz = voz._CONF
        voz._CONF = self.conf
        personalidades._estado["original_voz"] = None
        personalidades._cfg["cfg"] = self.conf

    def tearDown(self):
        import memoria
        memoria.DB_PATH = self._db
        voz._CONF = self._voz
        self.P._estado["cache"] = None
        self.P._estado["original_voz"] = None

    def test_se_reconoce_como_lo_dirias(self):
        b = self.P.buscar
        self.assertEqual(b("ponte en modo mirrey"), "mirrey")
        self.assertEqual(b("háblame como abuelita"), "abuelita")
        self.assertEqual(b("personalidad de godín"), "godin")
        self.assertEqual(b("vuelve a ser normal"), "jarvis")
        self.assertIsNone(b("activa el modo realidad aumentada"))

    def test_se_queda_para_siempre_y_cambia_la_voz(self):
        self.P.cambiar_personalidad("abuelita")
        self.assertEqual(self.conf["edge_voz"], "es-MX-DaliaNeural")
        self.P._estado["cache"] = None                    # como si Jarvis se reiniciara
        self.assertEqual(self.P.actual()[0], "abuelita")  # sigue guardada
        self.assertIn("Abuelita", self.P.prompt())
        self.assertEqual(self.P.prompt(en_exposicion=True), "")   # al exponer, la clásica
        self.P.cambiar_personalidad("normal")
        self.assertEqual(self.conf, {"edge_voz": "es-MX-JorgeNeural", "edge_velocidad": "+5%",
                                     "edge_tono": "+0Hz"})   # su voz original de vuelta
        self.assertEqual(self.P.prompt(), "")

    def test_tambien_al_exponer_si_lo_pide(self):
        self.P.cambiar_personalidad("coach", tambien_en_exposicion=True)
        self.assertIn("Coach", self.P.prompt(en_exposicion=True))

    def test_todas_completas(self):
        for clave, p in self.P.PERSONALIDADES.items():
            for campo in ("nombre", "alias", "descripcion", "estilo", "ejemplos", "voz", "saludo"):
                self.assertIn(campo, p, clave)
            if clave != "jarvis":
                self.assertTrue(p["estilo"] and p["ejemplos"], clave)
            self.assertEqual(self.P.buscar(p["nombre"]), clave, clave)

    def test_atajo_para_cambiar(self):
        self.assertEqual(genesis.atajo_personalidad("Jarvis, ponte en modo mirrey"),
                         ("cambiar_personalidad", {"personalidad": "mirrey"}))
        self.assertEqual(genesis.atajo_personalidad("cambia tu personalidad a norteño"),
                         ("cambiar_personalidad", {"personalidad": "norteno"}))
        self.assertIsNone(genesis.atajo_personalidad("activa el modo realidad aumentada"))
        self.assertIsNone(genesis.atajo_personalidad("mi abuelita hace buen mole"))


class Habitos(unittest.TestCase):
    def setUp(self):
        import habitos
        self.H = habitos
        self._db = habitos.DB
        habitos.DB = Path(tempfile.mkdtemp()) / "habitos.db"
        habitos._estado.update(pendiente=None, ultima=0.0)

    def tearDown(self):
        self.H.DB = self._db
        self.H._estado.update(pendiente=None, ofrecer=None, libre=None)

    @staticmethod
    def _obs(dias_atras, hora, app, musica, titulo="Lofi Girl", fuente="Edge", base=None):
        base = base or time.time()
        d = datetime.datetime.fromtimestamp(base - dias_atras * 86400).replace(
            hour=int(hora), minute=int(round((hora % 1) * 60)) % 60, second=0)
        return {"ts": d.timestamp(), "dia": d.weekday(), "hora": hora, "app": app, "musica": musica,
                "fuente": fuente if musica else "", "titulo": titulo if musica else ""}

    def test_musica_de_fondo_vs_ver_un_video(self):
        sonando = [("Edge", "Lofi Girl", "", 4)]
        self.assertEqual(self.H.musica_de_fondo("code.exe", sonando)[0], 1)   # programas, suena
        self.assertEqual(self.H.musica_de_fondo("msedge.exe", sonando)[0], 0)  # lo estás viendo
        self.assertEqual(self.H.musica_de_fondo("msedge.exe", [("Spotify", "x", "", 4)])[0], 1)
        self.assertEqual(self.H.musica_de_fondo("code.exe", [("Edge", "x", "", 5)])[0], 0)  # en pausa

    def test_ofrece_musica_si_es_tu_costumbre(self):
        ahora = datetime.datetime.now().replace(hour=16, minute=10).timestamp()
        mismos = [d for d in range(1, 15)
                  if self.H._finde(self._obs(d, 16, "x", 0, base=ahora)["dia"])
                  == self.H._finde(datetime.datetime.fromtimestamp(ahora).weekday())][:5]
        historial = [self._obs(d, 16 + m / 60, "code.exe", 1, base=ahora) for d in mismos for m in range(0, 30, 3)]
        historial += [self._obs(d, 16 + m / 60, "code.exe", 0, base=ahora) for d in mismos for m in range(30, 40, 3)]
        recientes = [self._obs(0, 16 + (10 - i) / 60, "code.exe", 0, base=ahora) for i in range(6)]
        r = self.H.evaluar_musica(historial, recientes)
        self.assertIsNotNone(r)
        self.assertEqual(r["titulo"], "Lofi Girl")
        self.assertGreaterEqual(r["dias"], 3)

    def test_no_ofrece_sin_costumbre_o_con_musica(self):
        ahora = datetime.datetime.now().replace(hour=16, minute=10).timestamp()
        recientes = [self._obs(0, 16 + (10 - i) / 60, "code.exe", 0, base=ahora) for i in range(6)]
        sin = [self._obs(d, 16 + m / 60, "code.exe", 0, base=ahora) for d in range(1, 6) for m in range(0, 30, 3)]
        self.assertIsNone(self.H.evaluar_musica(sin, recientes))           # nunca pones música ahí
        con = [self._obs(d, 16 + m / 60, "code.exe", 1, base=ahora) for d in range(1, 6) for m in range(0, 30, 3)]
        sonando = [self._obs(0, 16 + (10 - i) / 60, "code.exe", 1, base=ahora) for i in range(6)]
        self.assertIsNone(self.H.evaluar_musica(con, sonando))             # ya tienes música
        pocos = [self._obs(1, 16 + m / 60, "code.exe", 1, base=ahora) for m in range(0, 58, 2)]
        self.assertIsNone(self.H.evaluar_musica(pocos, recientes))         # un solo día: no es hábito

    def test_rutina_a_la_misma_hora(self):
        ahora = datetime.datetime.now().replace(hour=9, minute=30).timestamp()
        hoy_finde = self.H._finde(datetime.datetime.fromtimestamp(ahora).weekday())
        acciones = []
        for d in range(1, 15):
            o = self._obs(d, 9.5, "", 0, base=ahora)
            if self.H._finde(o["dia"]) == hoy_finde and len(acciones) < 4:
                acciones.append({"ts": o["ts"], "dia": o["dia"], "hora": o["hora"],
                                 "clave": "abrir_app:spotify", "skill": "abrir_app",
                                 "args": '{"nombre": "Spotify"}', "app": ""})
        r = self.H.evaluar_rutinas(acciones, ahora, set(), set())
        self.assertEqual(r["skill"], "abrir_app")
        self.assertEqual(r["args"], {"nombre": "Spotify"})
        self.assertIsNone(self.H.evaluar_rutinas(acciones, ahora, {"abrir_app:spotify"}, set()))  # ya lo hiciste hoy
        lejos = datetime.datetime.now().replace(hour=18, minute=0).timestamp()
        self.assertIsNone(self.H.evaluar_rutinas(acciones, lejos, set(), set()))   # no es la hora

    def test_respuesta_si_hace_la_accion_y_no_cuenta_rechazo(self):
        dicho = []
        self.H._estado.update(ofrecer=lambda t: dicho.append(t) or True, libre=lambda: True)
        self.assertTrue(self.H._ofrecer("musica", "code.exe", "youtube", {"consulta": "lofi"}, "¿Música?"))
        self.assertIsNotNone(self.H.pendiente())
        self.assertEqual(self.H.responder(True), ("youtube", {"consulta": "lofi"}))
        self.H._ofrecer("musica", "code.exe", "youtube", {"consulta": "lofi"}, "¿Música?")
        self.assertIsNone(self.H.responder(False))
        ofertas = [dict(o) for o in self.H._q("SELECT * FROM ofertas")]
        self.assertEqual([o["respuesta"] for o in ofertas], ["si", "no"])
        self.assertEqual(self.H.rechazos(ofertas, "musica", "code.exe", time.time()), 1)

    def test_aprende_lo_que_pides(self):
        self.H.registrar_accion("abrir_app", {"nombre": "Spotify"}, "Abriendo Spotify.")
        self.H.registrar_accion("volumen", {"nivel": 50}, "ok")                    # no es rutina
        self.H.registrar_accion("abrir_app", {"nombre": "X"}, skills.Fallo("no"))  # falló
        filas = self.H._q("SELECT clave FROM acciones")
        self.assertEqual([f["clave"] for f in filas], ["abrir_app:spotify"])

    def test_resumen_con_pocos_datos(self):
        self.assertIn("aprendiendo", self.H.resumen())


class AnalizarSituaciones(unittest.TestCase):
    def test_platica_vs_orden(self):
        import cognicion
        self.assertTrue(cognicion.es_situacion("fíjate que mi jefe me dijo que me van a cambiar de área"))
        self.assertTrue(cognicion.es_situacion("estoy estresado con la escuela"))
        self.assertFalse(cognicion.es_situacion("abre spotify"))
        self.assertEqual(cognicion.nivel("te cuento: me ofrecieron un trabajo en Monterrey"), "profundo")
        self.assertIn("PLATICANDO", cognicion.reglas("profundo", "no sé qué hacer con mi novia"))
        self.assertNotIn("PLATICANDO", cognicion.reglas("rapido", "abre spotify"))


class MicrofonoSilenciado(unittest.TestCase):
    """Pasó: el micrófono quedó silenciado en Windows y Jarvis 'dejó de oír' sin decir nada."""

    def test_al_arrancar_lo_reactiva_y_avisa(self):
        orig = (escuchar.microfono_silenciado, escuchar.activar_microfono)
        estado = {"mudo": True}
        escuchar.microfono_silenciado = lambda: estado["mudo"]
        escuchar.activar_microfono = lambda: estado.update(mudo=False) or True
        avisos = []
        try:
            escuchar.vigilar_microfono(avisos.append)
            fin = time.time() + 3
            while not avisos and time.time() < fin:
                time.sleep(0.05)
        finally:
            escuchar.microfono_silenciado, escuchar.activar_microfono = orig
        self.assertFalse(estado["mudo"])
        self.assertIn("silenciado", avisos[0])

    def test_silencio_digital(self):
        guardado = list(escuchar.NIVELES)
        try:
            escuchar.NIVELES.clear()
            escuchar.NIVELES.extend([0.0] * 100)
            self.assertTrue(escuchar._sin_senal())
            escuchar.NIVELES.extend([0.003] * 30)   # el ruido normal de un cuarto
            self.assertFalse(escuchar._sin_senal())
        finally:
            escuchar.NIVELES.clear()
            escuchar.NIVELES.extend(guardado)


class CicloDelDia(unittest.TestCase):
    """Resumen de la mañana, pendientes y avisos de noche (sin internet ni cámara)."""

    def setUp(self):
        import ciclo
        import memoria
        self.C = ciclo
        self._db = memoria.DB_PATH
        memoria.DB_PATH = Path(tempfile.mkdtemp()) / "prueba.db"
        self._estado = dict(ciclo._estado)
        ciclo._estado.update(cfg={"expositor": {"presentador": "Abraham Torres"}}, ofrecer=None,
                             libre=None, noche=None)

    def tearDown(self):
        import memoria
        memoria.DB_PATH = self._db
        self.C._estado.clear()
        self.C._estado.update(self._estado)

    def test_pendientes(self):
        ayer = (datetime.date.today() - datetime.timedelta(days=1)).isoformat()
        self.C.anotar_pendiente("terminar el reporte", para=ayer)
        self.C.anotar_pendiente("pagar la luz", para="mañana")
        self.assertEqual([p["texto"] for p in self.C.pendientes_para()], ["terminar el reporte"])
        self.assertIn("reporte", self.C.ver_pendientes())
        self.assertIn("Tachado", self.C.completar_pendiente("ya terminé el reporte"))
        self.assertEqual(self.C.pendientes_para(), [])
        manana = datetime.date.today() + datetime.timedelta(days=1)
        self.assertEqual([p["texto"] for p in self.C.pendientes_para(manana)], ["pagar la luz"])

    def test_fechas_dichas(self):
        hoy = datetime.date(2026, 10, 6)  # martes
        self.assertEqual(self.C._fecha("mañana", hoy), datetime.date(2026, 10, 7))
        self.assertEqual(self.C._fecha("hoy", hoy), hoy)
        self.assertEqual(self.C._fecha("el viernes", hoy), datetime.date(2026, 10, 9))
        self.assertEqual(self.C._fecha("martes", hoy), datetime.date(2026, 10, 13))  # el próximo

    def test_resumen_completo_y_en_orden(self):
        ayer = (datetime.date.today() - datetime.timedelta(days=1)).isoformat()
        self.C.anotar_pendiente("terminar el reporte", para=ayer)
        texto = self.C.armar_resumen(datetime.datetime.now().replace(hour=8, minute=30),
                                     clima_fn=lambda: "En Cuautitlán hace 16 grados.",
                                     whatsapp_fn=lambda: "En WhatsApp tienes 1 chat sin responder: Mamá.",
                                     correo_fn=lambda: "En tu correo hay 1 importante: de salud, Hospital.")
        self.assertTrue(texto.startswith("Buenos días, Abraham. Son las 8:30 de la mañana."))
        orden = [texto.index(x) for x in ("16 grados", "terminar el reporte", "Mamá", "Hospital")]
        self.assertEqual(orden, sorted(orden))

    def test_resumen_aunque_falle_algo(self):
        def falla():
            raise OSError("sin internet")
        texto = self.C.armar_resumen(datetime.datetime.now().replace(hour=15), clima_fn=falla,
                                     whatsapp_fn=lambda: "", correo_fn=lambda: "")
        self.assertTrue(texto.startswith("Buenas tardes"))

    def test_una_vez_al_dia_y_no_de_noche(self):
        dia = datetime.datetime.now().replace(hour=9, minute=0)
        self.assertTrue(self.C.resumen_pendiente_hoy(dia))
        self.C._ajuste("resumen", self.C.dia_logico(dia).isoformat())
        self.assertFalse(self.C.resumen_pendiente_hoy(dia))
        self.assertFalse(self.C.resumen_pendiente_hoy(dia.replace(hour=23, minute=30)))
        self.assertTrue(self.C.es_de_noche(dia.replace(hour=23, minute=30)))
        self.assertTrue(self.C.es_de_noche(dia.replace(hour=2)))
        self.assertFalse(self.C.es_de_noche(dia.replace(hour=14)))

    def test_avisos_de_noche(self):
        import cansancio
        dicho = []
        self.C._estado.update(ofrecer=dicho.append, libre=lambda: True)
        orig = (cansancio.encender, cansancio.estado, self.C._inactivo_seg)
        cansancio.encender = lambda activo=True: None
        estado = {"cansado": False, "motivo": "", "perclos": 0.0, "bostezos": 0, "muestras": 50}
        cansancio.estado = lambda: dict(estado)
        self.C._inactivo_seg = lambda: 5
        try:
            noche = datetime.datetime.now().replace(hour=23, minute=10)
            self.C.revisar(noche)
            self.assertEqual(len(dicho), 1)
            self.assertIn("Son las 11:10 de la noche", dicho[0])
            self.C.revisar(noche)                       # no repite enseguida
            self.assertEqual(len(dicho), 1)
            self.C._estado["ultimo_aviso_noche"] -= 15 * 60
            estado.update(cansado=True, motivo="los ojos se te cierran seguido")
            self.C.revisar(noche)                       # pero si te ve cansado, sí
            self.assertEqual(len(dicho), 2)
            self.assertIn("te noto cansado", dicho[1])
        finally:
            cansancio.encender, cansancio.estado, self.C._inactivo_seg = orig


class CansancioPorCamara(unittest.TestCase):
    def test_ojos_cerrados_y_bostezos(self):
        import cansancio
        c = cansancio.Cansancio()
        t = 1000.0
        for i in range(120):                 # despierto: parpadeos normales
            c.agregar(t + i * 0.25, 0.9 if i % 20 == 0 else 0.1, 0.1, 0.05)
        self.assertFalse(c.estado(t + 30)["cansado"])
        c2 = cansancio.Cansancio()
        for i in range(120):                 # ojos cerrados un 30 % del tiempo
            cerrado = i % 10 < 3
            c2.agregar(t + i * 0.25, 0.8 if cerrado else 0.1, 0.8 if cerrado else 0.1, 0.05)
        est = c2.estado(t + 30)
        self.assertTrue(est["cansado"])
        self.assertIn("ojos", est["motivo"])
        c3 = cansancio.Cansancio()
        for bostezo in (0, 60):              # dos bostezos de 2 s en un minuto
            for i in range(10):
                c3.agregar(t + bostezo + i * 0.25, 0.1, 0.1, 0.8 if i < 8 else 0.1)
        self.assertEqual(c3.estado(t + 70)["bostezos"], 2)
        self.assertTrue(c3.estado(t + 70)["cansado"])


class CorreoImportante(unittest.TestCase):
    def test_de_mas_a_menos_importante(self):
        import correo
        ejemplos = [("Amazon <x@amazon.com.mx>", "Tu pedido fue enviado"),
                    ("OCC Mundial <alertas@occ.com.mx>", "Nuevas vacantes"),
                    ("GBM <noreply@gbm.com>", "Rendimiento de tu portafolio"),
                    ("BBVA <avisos@bbva.mx>", "Tu estado de cuenta"),
                    ("BBVA <alertas@bbva.mx>", "Alerta de seguridad: cargo no reconocido"),
                    ("Hospital Ángeles <citas@hospitalesangeles.com>", "Tu cita médica"),
                    ("Computrabajo <x@computrabajo.com>", "Te invitaron a una entrevista")]
        cs = [{"de": d, "asunto": a, "ts": time.time() - i, "leido": False} for i, (d, a) in enumerate(ejemplos)]
        orden = [(c["categoria"], c["asunto"]) for c in correo.ordenar(cs)]
        self.assertEqual([o[0] for o in orden], ["salud", "banco", "banco", "inversiones", "empleo", "empleo"])
        self.assertIn("Alerta", orden[1][1])        # lo urgente primero
        self.assertIn("entrevista", orden[4][1])
        self.assertNotIn("Amazon", " ".join(a for _c, a in orden))


class ConectarGmail(unittest.TestCase):
    """Pasó: Google rechazó la conexión (se pegó la contraseña normal, no la de aplicación)."""

    def test_valida_antes_de_intentar(self):
        import correo
        intentos = []
        guardados = []
        orig = correo.guardar
        correo.guardar = lambda c, k: guardados.append((c, k))
        try:
            ok, aviso = correo.conectar("yo@gmail.com", "MiContraseña123", probar_fn=intentos.append)
            self.assertFalse(ok)
            self.assertIn("16 letras", aviso)
            self.assertEqual(intentos, [])                  # ni lo intentó con Google
            ok, aviso = correo.conectar(" Abraham.T082 ", "abcd efgh ijkl mnop",
                                        probar_fn=lambda c, k: intentos.append((c, k)))
            self.assertTrue(ok)
            self.assertEqual(intentos, [("abraham.t082@gmail.com", "abcdefghijklmnop")])
            self.assertEqual(guardados, [("abraham.t082@gmail.com", "abcdefghijklmnop")])
        finally:
            correo.guardar = orig

    def test_rechazo_de_google_se_explica(self):
        import correo

        def rechaza(c, k):
            raise correo.imaplib.IMAP4.error("[AUTHENTICATIONFAILED] Invalid credentials")
        ok, aviso = correo.conectar("yo@gmail.com", "abcdefghijklmnop", probar_fn=rechaza)
        self.assertFalse(ok)
        self.assertIn("no aceptó", aviso)


class WhatsAppSinResponder(unittest.TestCase):
    """Con el formato real de la app de escritorio (WhatsApp para Windows 2.26)."""

    def test_filas_reales(self):
        import whatsapp
        f = whatsapp.interpretar_fila
        c = f("Mamá 10:32 AM 2 unread messages", "2 unread messages ¿Ya llegaste?")
        self.assertEqual((c["nombre"], c["hora"], c["no_leidos"], c["mensaje"], c["grupo"]),
                         ("Mamá", "10:32 AM", 2, "¿Ya llegaste?", False))
        g = f("Equipo 9:15 AM 3 unread messages", "3 unread messages Juan Pérez: ya subí el reporte")
        self.assertEqual((g["grupo"], g["remitente"], g["mensaje"]), (True, "Juan Pérez", "ya subí el reporte"))
        t = f("Ventas 9:50 AM 31 unread messages", "31 unread messages Ventas ~ Edgar : Photo")
        self.assertEqual((t["grupo"], t["remitente"]), (True, "Edgar"))
        m = f("Sistemas 8:00 AM 4 unread messages Sistemas Muted chat", "4 unread messages hola")
        self.assertTrue(m["silenciado"])
        self.assertEqual(f("Comunidad Yesterday 690 unread messages", "690 unread messages Comunidad")["mensaje"], "")
        self.assertEqual(f("Mamá 🌸 10:32 AM", "hola")["nombre"], "Mamá")   # sin emojis para la voz
        self.assertIsNone(f("Sin hora", "texto"))

    def test_resumen_personas_primero_y_grupos_juntos(self):
        import whatsapp
        f = whatsapp.interpretar_fila
        chats = [f("Ventas 9:50 AM 690 unread messages", "690 unread messages Ventas"),
                 f("Mamá 10:32 AM 2 unread messages", "2 unread messages ¿Ya llegaste?"),
                 f("Ventas 9:40 AM 107 unread messages", "107 unread messages Ventas ~ Edgar : Photo"),
                 f("Juan Sunday 1 unread message", "1 unread message viejo"),          # más de 24 h
                 f("Sistemas 8:00 AM 4 unread messages Muted chat", "4 unread messages x"),  # silenciado
                 f("Pepe 8:00 AM", "ya leído")]                                        # sin pendientes
        texto = whatsapp.redactar(chats)
        self.assertTrue(texto.startswith("En WhatsApp te escribieron 1 persona y no has respondido: "
                                         "Mamá, 2 mensajes: ¿Ya llegaste?"))
        self.assertIn("Ventas con 797", texto)
        self.assertNotIn("Juan", texto)
        self.assertNotIn("?.", texto)
        self.assertIn("1 grupo silenciado", texto)


class ClimaDicho(unittest.TestCase):
    def test_describe_con_consejos(self):
        import clima
        d = {"current": {"time": "2026-10-06T09:00", "temperature_2m": 16.4, "apparent_temperature": 16.0,
                         "weather_code": 2},
             "hourly": {"time": ["2026-10-06T10:00", "2026-10-06T16:00"], "precipitation_probability": [20, 80]},
             "daily": {"time": ["2026-10-06", "2026-10-07"], "weather_code": [61, 3],
                       "temperature_2m_max": [20.1, 21], "temperature_2m_min": [8.2, 9],
                       "precipitation_probability_max": [80, 10], "uv_index_max": [5, 4]}}
        texto = clima.describir(d, "Cuautitlán")
        self.assertIn("En Cuautitlán hace 16 grados", texto)
        self.assertIn("80% de probabilidad de lluvia como a las 4 de la tarde", texto)
        self.assertIn("paraguas", texto)
        self.assertIn("abrígate", texto)
        self.assertIn("Mañana: 21 grados", texto)


class MemoriaPorSignificado(unittest.TestCase):
    """semantica.py con un modelo de embeddings simulado (sin Ollama ni internet)."""

    TEMAS = {"viaje": ("viaje", "vacaciones", "oaxaca", "playa"), "mascota": ("mascota", "perro", "rocky", "gato"),
             "examen": ("examen", "prueba", "calculo", "parcial")}

    def _vec(self, texto):
        t = skills._norm(texto)
        v = np.array([sum(p in t for p in palabras) for palabras in self.TEMAS.values()] + [0.1], dtype=np.float32)
        return v / np.linalg.norm(v)

    def setUp(self):
        import memoria
        import nube
        import semantica
        self.S, self.M, self.N = semantica, memoria, nube
        self._db = memoria.DB_PATH
        memoria.DB_PATH = Path(tempfile.mkdtemp()) / "prueba.db"
        self._orig = (semantica.embeddings, nube.lista, nube.peticion)
        semantica.embeddings = lambda textos: np.vstack([self._vec(t) for t in textos])
        self.peticiones = []
        nube.lista = lambda: True
        nube.peticion = lambda metodo, ruta, **kw: self.peticiones.append((metodo, ruta, kw)) or \
            type("R", (), {"json": lambda s: []})()
        semantica._estado.update(cambios=0, version=-1, matriz=None)

    def tearDown(self):
        self.M.DB_PATH = self._db
        self.S.embeddings, self.N.lista, self.N.peticion = self._orig
        self.S._estado.update(cambios=0, version=-1, matriz=None)

    def test_encuentra_por_significado_y_respeta_el_olvido(self):
        self.M.guardar_mensaje("user", "Me voy de vacaciones a Oaxaca en diciembre")
        self.M.guardar_mensaje("assistant", "¡Qué buen plan!")
        self.M.guardar_mensaje("user", "Mi perro se llama Rocky")
        self.M.guardar_mensaje("user", "ok")                     # muy corto: no se indexa
        self.M.agregar_hecho("Tiene examen de cálculo el viernes")
        agregados, _ = self.S.indexar()
        self.assertEqual(agregados, 3)
        r = self.S.buscar("lo que te dije del viaje")
        self.assertIn("Oaxaca", r[0]["texto"])
        self.assertEqual(self.S.buscar("¿cuándo es mi prueba?")[0]["tipo"], "hecho")
        # subir a la nube
        self.assertEqual(self.S.subir(), 3)
        self.assertTrue(any(m == "POST" and "/rest/v1/recuerdos" in ruta for m, ruta, _ in self.peticiones))
        # olvidar la conversación del viaje: se borra del índice y de Supabase
        ids = self.M.mensajes_sobre("vacaciones Oaxaca")
        self.assertEqual(len(ids), 2)                             # tu mensaje y la respuesta
        self.M.borrar_mensajes(ids)
        _, borrados = self.S.indexar()
        self.assertEqual(borrados, 1)
        self.assertEqual(self.S.buscar("lo que te dije del viaje"), [])
        self.peticiones.clear()
        self.S.subir()
        self.assertTrue(any(m == "DELETE" and "recuerdos?id=in." in ruta for m, ruta, _ in self.peticiones))
        self.assertEqual(self.S._q("SELECT id FROM borrar_en_nube"), [])

    def test_si_la_nube_falla_el_borrado_queda_pendiente(self):
        self.M.guardar_mensaje("user", "Mi perro se llama Rocky y es café")
        self.S.indexar()
        self.S.subir()

        def falla(metodo, ruta, **kw):
            raise self.N.NubeError("sin red")
        self.M.borrar_mensajes(self.M.mensajes_sobre("perro Rocky"))
        self.S.indexar()
        self.N.peticion = falla
        with self.assertRaises(self.N.NubeError):
            self.S.subir()
        self.assertEqual(len(self.S._q("SELECT id FROM borrar_en_nube")), 1)   # no se pierde

    def test_recordar_conversacion_por_significado(self):
        self.M.guardar_mensaje("user", "Me voy de vacaciones a Oaxaca en diciembre")
        self.M.guardar_mensaje("assistant", "Suena increíble")
        self.S.indexar()
        self.S._estado["hilo"] = "prueba"   # disponible
        try:
            texto = self.M.recordar_conversacion("el viaje")
        finally:
            self.S._estado["hilo"] = None
        self.assertIn("Oaxaca", texto)


class RespaldosCifrados(unittest.TestCase):
    def test_cifrado(self):
        import respaldo
        blob = respaldo.cifrar(b"mis datos", "frase secreta larga")
        self.assertNotIn(b"mis datos", blob)
        self.assertEqual(respaldo.descifrar(blob, "frase secreta larga"), b"mis datos")
        with self.assertRaises(ValueError):
            respaldo.descifrar(blob, "otra frase")

    def test_zip_sin_claves_y_restauracion(self):
        import io
        import sqlite3
        import zipfile
        import nube
        import respaldo
        tmp = Path(tempfile.mkdtemp())
        datos = tmp / "datos"
        datos.mkdir()
        con = sqlite3.connect(datos / "genesis.db")
        con.execute("CREATE TABLE t (x)")
        con.execute("INSERT INTO t VALUES ('hola')")
        con.commit()
        con.close()
        (datos / "gmail.json").write_text("secreto")            # no debe ir
        (datos / "apps.json").write_text("{}")
        avatares = tmp / "Avatares" / "Vault Boy"
        avatares.mkdir(parents=True)
        (avatares / "saludo.gif").write_bytes(b"GIF89a")
        (tmp / "config.json").write_text('{"name": "Jarvis"}')
        z = respaldo.armar_zip(datos, tmp / "Avatares", tmp / "config.json")
        nombres = zipfile.ZipFile(io.BytesIO(z)).namelist()
        self.assertIn("datos/genesis.db", nombres)
        self.assertIn("Avatares/Vault Boy/saludo.gif", nombres)
        self.assertNotIn("datos/gmail.json", nombres)
        blob = respaldo.cifrar(z, "frase de prueba")
        # restaurar en otra carpeta (como en otra PC)
        destino = Path(tempfile.mkdtemp())
        orig = (nube.peticion, respaldo.listar, respaldo.AVATARES)
        nube.peticion = lambda *a, **k: type("R", (), {"content": blob})()
        respaldo.listar = lambda equipo=None: [{"name": "2026-10-06_0300.jarvis"}]
        respaldo.AVATARES = destino / "Avatares"
        try:
            respaldo.restaurar("frase de prueba", destino=destino)
        finally:
            nube.peticion, respaldo.listar, respaldo.AVATARES = orig
        con = sqlite3.connect(destino / "datos" / "genesis.db")
        self.assertEqual(con.execute("SELECT x FROM t").fetchone()[0], "hola")
        con.close()
        self.assertTrue((destino / "Avatares" / "Vault Boy" / "saludo.gif").exists())


class ConectarSupabase(unittest.TestCase):
    def test_validaciones(self):
        import nube
        self.assertEqual(nube.normalizar_url("abcdefghijklmnopqrst"), "https://abcdefghijklmnopqrst.supabase.co")
        self.assertEqual(nube.normalizar_url("https://abcdefghijklmnopqrst.supabase.co/rest/v1/"),
                         "https://abcdefghijklmnopqrst.supabase.co")
        ok, aviso = nube.conectar("https://abcdefghijklmnopqrst.supabase.co", "sb_publishable_xxx",
                                  "frase larga", abrir_editor=False)
        self.assertFalse(ok)
        self.assertIn("pública", aviso)
        ok, aviso = nube.conectar("https://ejemplo.com", "sb_secret_xxx", "frase larga", abrir_editor=False)
        self.assertFalse(ok)
        ok, aviso = nube.conectar("https://abcdefghijklmnopqrst.supabase.co", "sb_secret_xxx", "corta",
                                  abrir_editor=False)
        self.assertFalse(ok)
        self.assertIn("8 caracteres", aviso)

    def test_esquema_seguro(self):
        sql = (Path(__file__).resolve().parents[1] / "supabase" / "esquema.sql").read_text(encoding="utf-8")
        self.assertIn("enable row level security", sql)
        self.assertIn("vector(768)", sql)
        self.assertIn("buscar_recuerdos", sql)
        self.assertNotIn("drop table", sql.lower())


class SupabaseAutomatico(unittest.TestCase):
    """configurar_automatico con la API de Supabase simulada."""

    def setUp(self):
        import nube
        self.N = nube
        self._orig = (nube._api, nube.estado_tabla, nube.guardar_json, nube.asegurar_bucket, nube.time.sleep)
        self.guardado = {}
        nube.guardar_json = lambda ruta, datos: self.guardado.update({Path(ruta).name: datos})
        nube.estado_tabla = lambda datos=None: "lista"
        nube.asegurar_bucket = lambda datos=None: None
        nube.time.sleep = lambda s: None
        self.llamadas = []

    def tearDown(self):
        (self.N._api, self.N.estado_tabla, self.N.guardar_json, self.N.asegurar_bucket,
         self.N.time.sleep) = self._orig
        self.N._cache.update(datos=None, lista=None)

    def _api_falsa(self, proyectos, crear_error=None):
        def api(token, metodo, ruta, **kw):
            self.llamadas.append((metodo, ruta))
            if ruta == "/projects" and metodo == "GET":
                return proyectos
            if ruta == "/organizations":
                return [{"id": "org1", "name": "abraham"}]
            if ruta == "/projects" and metodo == "POST":
                if crear_error:
                    raise self.N.NubeError(crear_error)
                return {"id": "nuevoref123456789012"}
            if ruta.endswith("/database/query"):
                return []
            if "api-keys" in ruta:
                return [{"name": "anon", "type": "publishable", "api_key": "sb_publishable_x"},
                        {"name": "default", "type": "secret", "api_key": "sb_secret_y"}]
            if ruta.startswith("/projects/"):
                return {"status": "ACTIVE_HEALTHY"}
        return api

    def test_crea_el_proyecto_y_guarda_solo_la_clave_del_proyecto(self):
        self.N._api = self._api_falsa([])
        ok, aviso = self.N.configurar_automatico("sbp_" + "a" * 40, "mi frase larga", avisar=lambda t: None)
        self.assertTrue(ok, aviso)
        self.assertIn(("POST", "/projects"), self.llamadas)
        datos = self.guardado["nube.json"]
        self.assertEqual(datos["url"], "https://nuevoref123456789012.supabase.co")
        self.assertEqual(datos["clave"], "sb_secret_y")          # la secreta, no la pública
        self.assertNotIn("sbp_", str(self.guardado))             # el token de la cuenta no se guarda

    def test_reutiliza_el_proyecto_jarvis(self):
        self.N._api = self._api_falsa([{"id": "yaexiste12345678901x", "name": "jarvis"}])
        ok, _ = self.N.configurar_automatico("sbp_" + "a" * 40, "mi frase larga", avisar=lambda t: None)
        self.assertTrue(ok)
        self.assertNotIn(("POST", "/projects"), self.llamadas)

    def test_limite_de_proyectos_gratis(self):
        self.N._api = self._api_falsa([], crear_error="402: The maximum limit of 2 free projects")
        ok, aviso = self.N.configurar_automatico("sbp_" + "a" * 40, "mi frase larga", avisar=lambda t: None)
        self.assertFalse(ok)
        self.assertIn("máximo", aviso)


class HonestidadSobreSusAcciones(unittest.TestCase):
    """Jarvis le dijo "claro, aquí ando jugando por ti" con Mortal Kombat sin mandar una tecla."""

    def setUp(self):
        import bitacora
        self.B = bitacora
        self._antes = list(bitacora._registro)
        bitacora._registro.clear()

    def tearDown(self):
        self.B._registro.clear()
        self.B._registro.extend(self._antes)

    def test_detecta_la_pregunta(self):
        for q in ("estás jugando por mí?", "Jarvis, ¿estás jugando por mí?", "¿fuiste tú el que movió el mouse?",
                  "¿tú tomaste el control?"):
            self.assertTrue(self.B.es_pregunta(q), q)
        for q in ("Pon música mientras estoy jugando", "¿Tienes el control de Spotify?", "abre Mortal Kombat"):
            self.assertFalse(self.B.es_pregunta(q), q)
        self.assertEqual(genesis.atajo_que_hice("¿estás jugando por mí?"), ("que_hice", {}))

    def test_contesta_con_lo_que_de_verdad_hizo(self):
        self.assertTrue(self.B.que_hice().startswith("No, no fui yo"))
        self.assertIn("NO ejecutaste", self.B.resumen())
        self.B.anotar("clima", {})                       # no mueve nada en la PC
        self.assertTrue(self.B.que_hice().startswith("No"))
        self.B.anotar("escribir_en", {"texto": "hola"})
        self.assertTrue(self.B.que_hice().startswith("Sí"))
        self.assertIn("escribir_en", self.B.resumen())
        # lo de hace más de 10 minutos ya no cuenta
        self.B._registro.clear()
        self.B._registro.append((time.time() - 3600, "escribir_en", {}, True))
        self.assertTrue(self.B.que_hice().startswith("No"))

    def test_etiquetas_mal_formadas_no_se_dicen(self):
        import acciones
        self.assertEqual(acciones.separar("Obvio, mi rey. [ACCION]pensando] [ACCION: saludo]"),
                         ("Obvio, mi rey.", "saludo"))
        self.assertEqual(acciones.separar("Listo [Acción - ejecutando]"), ("Listo", "ejecutando"))
        self.assertEqual(acciones.separar("hola [ACCION]")[0], "hola")


class JarvisJuega(unittest.TestCase):
    """juegos.py: botones de verdad (control virtual / teclado) y combos por voz."""

    def setUp(self):
        import juegos
        self.J = juegos
        self._orig = (juegos.ARCHIVO, juegos._juego_al_frente)
        juegos.ARCHIVO = Path(tempfile.mkdtemp()) / "juegos.json"

    def tearDown(self):
        self.J.ARCHIVO, self.J._juego_al_frente = self._orig

    def test_notacion_en_espanol_y_en_siglas(self):
        I = self.J.interpretar
        self.assertEqual(I("atrás adelante dos"), [["B"], ["F"], ["2"]])
        self.assertEqual(I("B, F, 2"), I("atrás, adelante y dos"))     # "y" no es "al mismo tiempo"
        self.assertEqual(I("abajo más bloqueo, cuatro"), [["D", "BL"], ["4"]])
        self.assertEqual(I("D+BL"), [["D", "BL"]])
        with self.assertRaises(ValueError):
            I("abre youtube")
        self.assertTrue(self.J.es_notacion("Jarvis, atrás adelante dos"))
        self.assertFalse(self.J.es_notacion("abre youtube"))
        self.assertFalse(self.J.es_notacion("abajo"))                  # una dirección sola no basta

    def test_adelante_y_atras_segun_el_lado(self):
        p = self.J.PERFILES["mortal_kombat"]
        self.assertEqual(self.J.resolver(["F", "1"], p, "izquierda"), ["DPAD_R", "X"])
        self.assertEqual(self.J.resolver(["F", "1"], p, "derecha"), ["DPAD_L", "X"])
        teclado = dict(p, control="teclado")
        self.assertEqual(self.J.resolver(["B", "BL"], teclado, "izquierda"), ["A", "SPACE"])

    def test_ejecuta_con_el_ritmo_del_juego(self):
        vistos = []
        p = dict(self.J.PERFILES["mortal_kombat"], cuadros=1, hueco=1)
        t0 = time.perf_counter()
        self.J.ejecutar(self.J.interpretar("B F 2 espera 4"), p, "izquierda",
                        poner=lambda e, abajo: vistos.append((tuple(e), abajo)))
        self.assertEqual([e for e, abajo in vistos if abajo], [("DPAD_L",), ("DPAD_R",), ("Y",), ("B",)])
        self.assertEqual(len([1 for _, abajo in vistos if not abajo]), 4)   # todo se suelta
        self.assertGreater(time.perf_counter() - t0, 0.2)                  # la espera sí espera

    def test_nunca_en_juegos_en_linea(self):
        self.assertTrue(self.J.bloqueado("cod.exe"))
        self.assertTrue(self.J.bloqueado("destiny2.exe"))
        self.assertTrue(self.J.bloqueado("VALORANT-Win64-Shipping.exe"))
        self.assertFalse(self.J.bloqueado("MK11.exe"))
        self.assertFalse(self.J.bloqueado("javaw.exe"))
        self.J._juego_al_frente = lambda: ("cod.exe", "Call of Duty")
        self.assertIsInstance(self.J.combo_juego("B F 2"), skills.Fallo)
        self.assertIsNone(genesis.atajo_juego("atrás adelante dos"))

    def test_combos_guardados_por_voz(self):
        self.J._juego_al_frente = lambda: ("mk11.exe", "Mortal Kombat 11")
        self.assertEqual(self.J.perfil_de("mk11.exe"), "mortal_kombat")
        self.assertIn("Gancho", self.J.guardar_combo("Gancho", "atrás adelante dos"))
        self.assertEqual(genesis.atajo_juego("Jarvis, haz el gancho"), ("combo_juego", {"secuencia": "gancho"}))
        self.assertEqual(genesis.atajo_juego("atrás adelante dos"), ("combo_juego", {"secuencia": "atras adelante dos"}))
        self.assertIsNone(genesis.atajo_juego("¿cómo te llamas?"))
        self.J._juego_al_frente = lambda: ("code.exe", "Visual Studio Code")   # sin juego al frente
        self.assertIsNone(genesis.atajo_juego("atrás adelante dos"))



class DetectorParaMi(unittest.TestCase):
    """para_mi.py: red local que decide si lo dicho sin llamarlo era para Jarvis."""

    def _vec(self, texto):
        t = skills._norm(texto)
        orden = sum(p in t for p in ("abre", "pon", "dime", "busca", "cuanto", "siguiente"))
        fondo = sum(p in t for p in ("video", "suscribete", "mama", "canal", "partido", "gol"))
        v = np.array([orden, fondo, 0.1], dtype=np.float32)
        v = v + np.random.RandomState(abs(hash(t)) % 2 ** 31).normal(0, 0.05, 3).astype(np.float32)
        return v / np.linalg.norm(v)

    def setUp(self):
        import para_mi
        self.P = para_mi
        carpeta = Path(tempfile.mkdtemp())
        self._orig = (para_mi.MODELO, para_mi.CORRECCIONES, para_mi.codificar)
        para_mi.MODELO, para_mi.CORRECCIONES = carpeta / "modelo.joblib", carpeta / "correcciones.jsonl"
        para_mi.codificar = lambda textos: np.vstack([self._vec(t) for t in textos])
        para_mi._estado["modelo"] = None
        para_mi._silenciados.clear()

    def tearDown(self):
        self.P.MODELO, self.P.CORRECCIONES, self.P.codificar = self._orig
        self.P._estado["modelo"] = None
        self.P._silenciados.clear()

    def test_quita_el_nombre(self):
        self.assertEqual(self.P.quitar_nombre("Oye Jarvis, abre Teams"), "abre Teams")
        self.assertEqual(self.P.quitar_nombre("abre Teams, Jarvis"), "abre Teams")

    def test_lee_el_registro(self):
        lineas = ["Tú: Jarvis, abre Teams por favor", "[Skill] teams_abrir {}",
                  "[Ignoro «suscríbete a mi canal y dale like»: venía de la computadora, no de ti]",
                  "Tú (sin llamarme): y el partido de ayer qué", "[No era para mí: me quedo callado]",
                  "Tú (sin llamarme): en el video se ve un carro", "Jarvis: Qué interesante.",
                  "Tú (sin llamarme): pásame la sal", "[La red dice que no era para mí: me quedo callado]"]
        si, no, dudosos = self.P.del_registro(lineas)
        self.assertEqual(si, ["abre Teams por favor"])
        self.assertEqual(no, ["suscríbete a mi canal y dale like", "y el partido de ayer qué"])
        self.assertEqual(dudosos, ["en el video se ve un carro"])   # lo que silencia la red no se reaprende

    def test_umbrales(self):
        probs = [0.01, 0.02, 0.03, 0.04, 0.05, 0.5, 0.97, 0.98, 0.99, 0.995, 0.999]
        ys = [0, 0, 0, 0, 0, 1, 1, 1, 1, 1, 1]
        self.assertEqual(self.P.umbrales(probs, ys), (0.05, 0.5))
        self.assertEqual(self.P.umbrales([0.1] * 10, [0, 1] * 5), (0.0, 1.0))   # nunca seguro

    def test_entrena_decide_y_aprende_de_sus_errores(self):
        si = ["abre el correo ahora", "pon música de Bad Bunny", "dime el clima de hoy", "busca recetas de pozole",
              "cuánto cuesta el dólar", "la siguiente canción porfa"] * 4
        no = ["en este video vamos a ver", "suscríbete a mi canal amigos", "mamá ya voy a comer",
              "qué golazo en el partido", "dale like al video", "gol gol gol del américa"] * 4
        m = self.P.entrenar(lambda *_: None, datos=(si, no, ["en el video se ve un carro"],
                                                     ["abre spotify", "pon la siguiente"], ["canal de cocina"]))
        self.assertGreaterEqual(m["exactitud"], 0.9)
        self.assertEqual(self.P.decidir("suscríbete al canal y mira el video", callarse=0.2), "no")
        self.assertNotEqual(self.P.decidir("abre el correo y dime qué hay", callarse=0.2), "no")
        self.assertIsNone(self.P.decidir("suscríbete al canal", callarse=0.0))   # 0 = apagado
        # la red silenció algo y lo repetiste llamándolo: corrección
        self.P._silenciados.append((time.time(), "pon el video de cocina"))
        self.assertTrue(self.P.llamado("Jarvis, pon el video de cocina"))
        self.assertEqual(self.P.correcciones(), ["pon el video de cocina"])
        self.assertFalse(self.P.llamado("Jarvis, qué hora es"))


class VisionConLimites(unittest.TestCase):
    """vision.py: una nube que llegó a su límite queda en pausa; lo de fondo va al modelo local."""

    def setUp(self):
        import vision
        self.V = vision
        self._orig = (vision._nube_varias, vision._local_varias, vision._proveedores, vision._usar_nube)
        self.llamadas = []
        vision._proveedores = lambda conf: [{"modelo": "nube-a", "url": "x"}, {"modelo": "nube-b", "url": "y"}]
        vision._usar_nube = lambda conf: bool(vision.libres(conf))
        vision._local_varias = lambda conf, s, i, l: self.llamadas.append("local") or "visto en local"
        cerebro._saturado.clear()

    def tearDown(self):
        self.V._nube_varias, self.V._local_varias, self.V._proveedores, self.V._usar_nube = self._orig
        cerebro._saturado.clear()

    def test_el_limite_pone_en_pausa(self):
        def nube(conf, n, s, i, l):
            self.llamadas.append(n["modelo"])
            if n["modelo"] == "nube-a":
                raise RuntimeError("Error code: 429 - Rate limit reached. Please try again in 1m30s")
            return "visto en b"
        self.V._nube_varias = nube
        self.assertEqual(self.V._ver({}, "s", "i", ["b64"]), "visto en b")
        self.assertGreater(cerebro._saturado["nube-a"] - time.time(), 80)    # pausa de 1m30s
        self.llamadas.clear()
        self.assertEqual(self.V._ver({}, "s", "i", ["b64"]), "visto en b")
        self.assertEqual(self.llamadas, ["nube-b"])                          # ni lo intenta con la "a"

    def test_todas_en_pausa_va_directo_al_local(self):
        self.V._nube_varias = lambda *a: self.fail("no debía llamar a la nube")
        cerebro._saturado.update({"nube-a": time.time() + 60, "nube-b": time.time() + 60})
        self.assertEqual(self.V._ver({}, "s", "i", ["b64"]), "visto en local")

    def test_lo_de_fondo_no_gasta_la_nube(self):
        self.V._nube_varias = lambda *a: self.fail("lo de fondo no debía usar la nube")
        self.assertEqual(self.V._ver({}, "s", "i", ["b64"], fondo=True), "visto en local")


class WhisperSinMemoria(unittest.TestCase):
    """Sin RAM, Whisper tumbaba a Jarvis completo (mkl_malloc: failed to allocate memory)."""

    def setUp(self):
        self._orig = (escuchar._cargar_modelo, escuchar._transcribir_nube, escuchar.liberar_memoria,
                      escuchar._avisar_memoria, escuchar._gpu["ok"], escuchar._cfg)
        self.liberadas = []
        escuchar.liberar_memoria = lambda: self.liberadas.append(1)
        escuchar._avisar_memoria = lambda: None
        escuchar._gpu["ok"] = False
        escuchar._cfg = lambda: {"stt": {"modo": "auto"}}

    def tearDown(self):
        (escuchar._cargar_modelo, escuchar._transcribir_nube, escuchar.liberar_memoria,
         escuchar._avisar_memoria, escuchar._gpu["ok"], escuchar._cfg) = self._orig

    def _sin_ram(self, nombre):
        raise RuntimeError("mkl_malloc: failed to allocate memory")

    def test_no_tumba_a_jarvis_y_rescata_con_la_nube(self):
        escuchar._cargar_modelo = self._sin_ram
        nube = iter([None, "abre spotify"])   # la 1a vez (nube=True) no contesta; la de rescate sí
        escuchar._transcribir_nube = lambda audio, prompt: next(nube)
        audio = np.zeros(16000, dtype=np.float32)
        self.assertEqual(escuchar.transcribir(audio, nube=True), "abre spotify")
        self.assertEqual(len(self.liberadas), 2)          # liberó memoria en cada intento
        escuchar._transcribir_nube = lambda audio, prompt: None
        self.assertEqual(escuchar.transcribir(audio, nube=True), "")   # sin nube: "no te escuché"

    def test_otros_errores_no_se_esconden(self):
        def roto(nombre):
            raise ValueError("otra cosa")
        escuchar._cargar_modelo = roto
        with self.assertRaises(ValueError):
            escuchar._transcribir_local(np.zeros(10, dtype=np.float32), "small", None)

    def test_reconoce_falta_de_memoria(self):
        self.assertTrue(escuchar.es_falta_de_memoria(MemoryError()))
        self.assertTrue(escuchar.es_falta_de_memoria(RuntimeError("CUDA failed: out of memory")))
        self.assertFalse(escuchar.es_falta_de_memoria(RuntimeError("cublas64_12.dll not found")))


class AprendeSoloDeTi(unittest.TestCase):
    """Guardaba "Se llama Jarvis" y "Su nombre es Abraham Jarvis" de una canción y un video."""

    def setUp(self):
        import para_mi
        self.P = para_mi
        self._orig = (para_mi.probabilidad, escuchar.frase_de_la_pc)
        escuchar.frase_de_la_pc = lambda minimo=0.5: False

    def tearDown(self):
        self.P.probabilidad, escuchar.frase_de_la_pc = self._orig

    def test_de_quien_aprende(self):
        self.P.probabilidad = lambda t: 0.3
        self.assertTrue(genesis._es_de_fiar("me gusta el rock", seguimiento=False))    # lo llamaste
        self.assertFalse(genesis._es_de_fiar("amiguitos, mi nombre es...", seguimiento=True))
        self.P.probabilidad = lambda t: 0.95
        self.assertTrue(genesis._es_de_fiar("estudio ingeniería en sistemas", seguimiento=True))
        escuchar.frase_de_la_pc = lambda minimo=0.5: True                            # sonaba la PC
        self.assertFalse(genesis._es_de_fiar("estudio ingeniería en sistemas", seguimiento=True))

    def test_no_guarda_hechos_sobre_jarvis(self):
        import cognicion
        import memoria
        guardados = []
        orig = (memoria.listar_hechos, memoria.agregar_hecho)
        memoria.listar_hechos = lambda: []
        memoria.agregar_hecho = guardados.append
        try:
            cognicion._guardar(["Se llama Jarvis", "Su nombre es Abraham Jarvis", "Le gusta el rock"])
        finally:
            memoria.listar_hechos, memoria.agregar_hecho = orig
        self.assertEqual(guardados, ["Le gusta el rock"])


class ControlDesdeElTelefono(unittest.TestCase):
    """remoto.py + genesis: órdenes desde el teléfono por Supabase."""

    def setUp(self):
        import nube
        import remoto
        self.R, self.N = remoto, nube
        self.ordenes = [{"id": 7, "texto": "¿qué hora es?", "creado": "x"}]
        self.respuestas, self.tomadas = [], []

        class Resp:
            def __init__(s, datos):
                s._d = datos

            def json(s):
                return s._d

        def peticion(metodo, ruta, json=None, encabezados=None, **kw):
            if ruta.startswith("/rest/v1/telefonos"):
                return Resp([{"id": 1}])
            if metodo == "GET" and "estado=eq.pendiente" in ruta:
                return Resp(self.ordenes)
            if metodo == "PATCH" and "estado=eq.pendiente" in ruta:
                self.tomadas.append(ruta)
                return Resp([{"id": 7}] if len(self.tomadas) == 1 else [])
            if metodo == "PATCH":
                self.respuestas.append(json)
            return Resp([])
        self._orig = (nube.peticion, remoto.entregar, dict(remoto._estado))
        nube.peticion = peticion
        remoto._estado.update(telefonos=None, t_latido=0.0, t_limpieza=0.0, procesando=set())

    def tearDown(self):
        self.N.peticion, self.R.entregar = self._orig[0], self._orig[1]
        self.R._estado.clear()
        self.R._estado.update(self._orig[2])

    def test_la_llave_va_despues_del_gato(self):
        d = self.R.direccion("LLAVE/+x", url="https://abc.supabase.co", publica="sb_publishable_1",
                             pagina="https://p.io/c/")
        pagina, fragmento = d.split("#", 1)
        self.assertEqual(pagina, "https://p.io/c/")             # al servidor solo llega esto
        self.assertIn("t=LLAVE%2F%2Bx", fragmento)
        self.assertEqual(len(self.R.huella("x")), 64)

    def test_toma_cada_orden_una_sola_vez(self):
        entregadas = []
        self.R.entregar = entregadas.append
        self.assertEqual(self.R.revisar_una_vez(), 1)
        self.assertEqual(self.R.revisar_una_vez(), 0)          # ya está en proceso
        self.R._estado["procesando"].clear()
        self.assertEqual(self.R.revisar_una_vez(), 0)          # Supabase dice que ya no está pendiente
        self.assertEqual([o["id"] for o in entregadas], [7])

    def test_genesis_responde_al_telefono_y_no_hace_lo_peligroso(self):
        orig = genesis._procesar
        try:
            def procesar(cfg, history, user, escrito, interruptor, **kw):
                self.assertFalse(cfg["voz_activa"])               # no habla en la PC
                genesis._entregar(cfg, type("T", (), {"decir": lambda s, t: None})(), "Son las 5.", [])
            genesis._procesar = procesar
            genesis._atender_remota({"voz_activa": True}, [], {"id": 7, "texto": "¿qué hora es?"}, None)
            self.assertEqual(self.respuestas[-1]["respuesta"], "Son las 5.")
            self.assertEqual(self.respuestas[-1]["estado"], "hecha")

            def procesar_peligroso(cfg, history, user, escrito, interruptor, **kw):
                r = genesis.ejecutar_herramienta(cfg, "apagar_equipo", {})
                self.assertIsInstance(r, skills.Fallo)
                genesis._entregar(cfg, type("T", (), {"decir": lambda s, t: None})(), "Falta de permisos.", [])
            genesis._procesar = procesar_peligroso
            genesis._atender_remota({}, [], {"id": 8, "texto": "apaga la compu"}, None)
            self.assertIn("Por seguridad no lo hice", self.respuestas[-1]["respuesta"])
        finally:
            genesis._procesar = orig
        self.assertIsNone(genesis._remota_actual["orden"])


class WhatsAppEscribirYLlamar(unittest.TestCase):
    """whatsapp_chat.py: elegir bien el chat antes de escribir nada."""

    def setUp(self):
        import whatsapp_chat
        self.W = whatsapp_chat

    def test_parecido_de_nombres(self):
        p = self.W.parecido
        self.assertEqual(p("Ana López", "Ana López"), 1.0)
        self.assertGreaterEqual(p("mi mamá", "Mamá ❤️"), 0.75)
        self.assertGreaterEqual(p("ana", "Ana López"), 0.75)
        self.assertLess(p("Ana", "Mariana Ruiz"), 0.75)
        self.assertLess(p("Luis", "Ana López"), 0.75)

    def test_resultados_por_secciones(self):
        filas = ["Chats", "Ana López 4:20 PM Nos vemos mañana", "Abraham Tc (You) Message yourself",
                 "Contacts", "Ana Sofía Hey there! I am using WhatsApp.", "Messages",
                 "Ana López 9/24/2026 ana me debe 200"]
        c = self.W.candidatos(filas)
        self.assertEqual([x[1] for x in c], ["Ana López", "Abraham Tc", "Ana Sofía Hey there!"])
        self.assertTrue(c[1][2])                                          # el chat propio
        self.assertEqual(self.W.elegir("Ana López", c), ("ok", c[0]))
        estado, nombres = self.W.elegir("Ana", c)
        self.assertEqual(estado, "varios")                                # Ana López y Ana Sofía: pregunta
        self.assertEqual(self.W.elegir("yo", c), ("ok", c[1]))
        self.assertEqual(self.W.elegir("Pedro", c), ("ninguno", None))

    def test_caja_del_mensaje(self):
        caja = type("C", (), {"element_info": type("I", (), {"name": "Type a message to Ana López"})()})()
        self.assertEqual(self.W.destinatario(caja), "Ana López")
        self.assertEqual(self.W.destinatario(None), "")

    def test_desde_el_telefono_no_se_confirma(self):
        genesis._remota_actual.update(orden={"id": 1}, bloqueada=None)
        try:
            self.assertFalse(genesis.confirmar({}, "¿Le mando a Ana: hola?"))
            self.assertTrue(genesis._remota_actual["bloqueada"])
        finally:
            genesis._remota_actual.update(orden=None, bloqueada=None)


class NombresQueSuenanIgual(unittest.TestCase):
    """Whisper escribió "jumcook" por "Yun Cook" y WhatsApp no lo encontraba."""

    def setUp(self):
        import whatsapp_chat
        self.W = whatsapp_chat
        self._agenda = whatsapp_chat.AGENDA
        whatsapp_chat.AGENDA = Path(tempfile.mkdtemp()) / "contactos.json"

    def tearDown(self):
        self.W.AGENDA = self._agenda

    def test_por_sonido(self):
        s = self.W.similitud
        for dicho, real in (("jumcook", "Yun Cook"), ("llun cuk", "Yun Cook"), ("jerardo", "Gerardo Peña"),
                            ("bictor", "Víctor"), ("jorge", "Yorge Ramírez"), ("mi mamá", "Mamá ❤️")):
            self.assertGreaterEqual(s(dicho, real), 0.95, (dicho, real))
        self.assertLess(s("Ana", "Mariana Ruiz"), 0.72)
        self.assertLess(s("Luis", "Ana López"), 0.72)

    def test_elige_al_mas_parecido_y_pregunta_si_empatan(self):
        c = [(1, "Yun Cook", False), (2, "Yuri Cortés", False), (3, "Abraham Tc", True)]
        self.assertEqual(self.W.elegir("jumcook", c), ("ok", c[0]))
        estado, nombres = self.W.elegir("Ana", [(1, "Ana López", False), (2, "Ana Sofía", False)])
        self.assertEqual(estado, "varios")
        self.assertEqual(self.W.elegir("Yuri Cortez", [(1, "Yuri Cortés", False)])[0], "ok")
        estado, sugerencia = self.W.elegir("Pedrito", [(1, "Pedro Infante", False)])
        self.assertEqual(estado, "ok" if self.W.similitud("Pedrito", "Pedro Infante") >= 0.72 else "ninguno")

    def test_rescate_y_agenda(self):
        self.assertEqual(self.W.consultas_de_rescate("jumcook")[:2], ["cook", "jumc"])
        self.W.recordar_nombres(["Yun Cook", "Mamá ❤️", "Yun Cook", "un chat"])
        self.assertEqual(self.W._agenda(), ["Yun Cook", "Mamá ❤️"])
        self.assertEqual(self.W.mas_parecido("jumcook", self.W._agenda())[0], "Yun Cook")
        self.assertIsNone(self.W.mas_parecido("Pedro", self.W._agenda()))


class GrietasDelRegistro(unittest.TestCase):
    """Errores reales que aparecieron en datos/genesis.log."""

    def test_un_error_de_cublas_no_tumba_a_jarvis(self):
        orig = (escuchar._cargar_modelo, escuchar._reprobar_gpu_luego, escuchar._gpu["ok"], escuchar._cfg)
        intentos = []

        class Modelo:
            def transcribe(self, audio, **kw):
                return [], None

        def cargar(nombre):
            intentos.append(escuchar._gpu["ok"])
            if escuchar._gpu["ok"]:
                raise RuntimeError("cuBLAS failed with status CUBLAS_STATUS_NOT_SUPPORTED")
            return Modelo()
        try:
            escuchar._cargar_modelo, escuchar._reprobar_gpu_luego = cargar, lambda: None
            escuchar._gpu["ok"] = True
            escuchar._cfg = lambda: {"stt": {"modo": "local"}}
            self.assertEqual(escuchar.transcribir(np.zeros(1600, dtype=np.float32)), "")
            self.assertEqual(intentos, [True, False])             # GPU falló -> procesador
            escuchar._cargar_modelo = lambda n: (_ for _ in ()).throw(ValueError("otra cosa"))
            self.assertEqual(escuchar.transcribir(np.zeros(1600, dtype=np.float32)), "")   # nunca truena
        finally:
            escuchar._cargar_modelo, escuchar._reprobar_gpu_luego, escuchar._gpu["ok"], escuchar._cfg = orig

    def test_el_resumen_dice_lo_que_no_pudo_revisar(self):
        import ciclo
        def falla():
            raise OSError("sin red")
        r = ciclo.armar_resumen(clima_fn=lambda: None, whatsapp_fn=falla, correo_fn=lambda: "")
        self.assertIn("No pude revisar el clima ni tu WhatsApp", r)
        self.assertNotIn("nada urgente", r)
        r = ciclo.armar_resumen(clima_fn=lambda: "", whatsapp_fn=lambda: "", correo_fn=lambda: "")
        self.assertNotIn("No pude revisar", r)

    def test_voz_sin_conexion_aparta_a_edge_de_inmediato(self):
        import voz
        orig = (voz._edge_trozos, voz._sapi_pcm, voz._motores, dict(voz._caidos), dict(voz._fallos))

        class ConnectionTimeoutError(Exception):
            pass

        def edge(frase):
            raise ConnectionTimeoutError("Connection timeout to host wss://speech.platform.bing.com")
            yield
        try:
            voz._edge_trozos = edge
            voz._sapi_pcm = lambda frase: np.zeros(100, dtype=np.int16)
            voz._motores = lambda vn, f=None: ["edge", "windows"]
            voz._caidos.clear()
            voz._fallos.clear()
            clip = voz._Clip("Hola, ¿cómo estás?")
            voz._generar(clip, None)
            self.assertGreater(voz._caidos.get("edge", 0), time.time())   # apartado al primer fallo
        finally:
            voz._edge_trozos, voz._sapi_pcm, voz._motores = orig[:3]
            voz._caidos.clear(); voz._caidos.update(orig[3])
            voz._fallos.clear(); voz._fallos.update(orig[4])
