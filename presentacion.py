"""PowerPoint por voz: abrir, iniciar, avanzar, regresar, saltar a una diapositiva, pantalla
negra, y saber qué dice cada diapositiva para poder explicarla al público.

Se controla por COM (la misma interfaz que usan las macros de Office), no con teclas: así da
igual qué ventana tenga el foco. Si una ventanita de Jarvis o una notificación se pone encima,
"siguiente" sigue funcionando. Si COM no está disponible (Office no instalado, otra app de
diapositivas, Google Slides en el navegador), se cae a mandar teclas a la ventana.
"""
import os
import re
import threading
import time

import pyautogui

import archivos
from skills import Callado, skill

EXTENSIONES = (".pptx", ".ppt", ".ppsx", ".pps", ".pptm", ".odp")
_contenido = {}      # ruta completa -> [ {numero, titulo, texto, notas}, ... ]
_lock = threading.Lock()

# Estados de SlideShowView.State (constantes de PowerPoint)
CORRIENDO, PAUSADO, NEGRA, BLANCA = 1, 2, 3, 4


# ---------- COM ----------
def _com():
    import pythoncom
    pythoncom.CoInitialize()  # cada hilo que usa COM tiene que iniciarlo; repetirlo no daña
    import win32com.client
    return win32com.client


def _app(crear=False):
    wc = _com()
    try:
        return wc.GetActiveObject("PowerPoint.Application")
    except Exception:
        if not crear:
            return None
    try:
        app = wc.Dispatch("PowerPoint.Application")
        app.Visible = True
        return app
    except Exception:
        return None


def _vista():
    """SlideShowView de la presentación que se está exponiendo, o None."""
    app = _app()
    try:
        if app is not None and app.SlideShowWindows.Count > 0:
            return app.SlideShowWindows.Item(1).View
    except Exception:
        pass
    return None


def _presentacion_activa():
    app = _app()
    if app is None:
        return None
    try:
        if app.SlideShowWindows.Count > 0:
            return app.SlideShowWindows.Item(1).Presentation
        if app.Presentations.Count > 0:
            return app.ActivePresentation
    except Exception:
        pass
    return None


def _ventana_clase(*clases):
    try:
        import win32gui
        for c in clases:
            h = win32gui.FindWindow(c, None)
            if h:
                return h
    except Exception:
        pass
    return None


def en_curso():
    """True si hay una presentación en modo pantalla completa."""
    try:
        if _vista() is not None:
            return True
    except Exception:
        pass
    return bool(_ventana_clase("screenClass"))


def reenfocar():
    """Devuelve el foco a la presentación (p. ej. tras cerrar una ventanita de confirmación),
    para que el clicker o las flechas del teclado sigan funcionando."""
    try:
        app = _app()
        if app is not None and app.SlideShowWindows.Count > 0:
            app.SlideShowWindows.Item(1).Activate()
    except Exception:
        pass


# ---------- Contenido de las diapositivas ----------
def _limpio(t):
    return re.sub(r"\s+", " ", str(t or "").replace("\r", " ").replace("\x0b", " ")).strip()


def _leer_com(pres):
    diapos = []
    for i in range(1, pres.Slides.Count + 1):
        s = pres.Slides.Item(i)
        titulo, textos, notas = "", [], ""
        try:
            if s.Shapes.HasTitle:
                titulo = _limpio(s.Shapes.Title.TextFrame.TextRange.Text)
        except Exception:
            pass
        for j in range(1, s.Shapes.Count + 1):
            try:
                sh = s.Shapes.Item(j)
                if sh.HasTextFrame and sh.TextFrame.HasText:
                    t = _limpio(sh.TextFrame.TextRange.Text)
                    if t and t != titulo:
                        textos.append(t)
            except Exception:
                continue
        try:
            for j in range(1, s.NotesPage.Shapes.Count + 1):
                sh = s.NotesPage.Shapes.Item(j)
                try:
                    if sh.PlaceholderFormat.Type == 2 and sh.HasTextFrame:  # 2 = cuerpo (notas)
                        notas = _limpio(sh.TextFrame.TextRange.Text)
                except Exception:
                    continue
        except Exception:
            pass
        diapos.append({"numero": i, "titulo": titulo, "texto": " / ".join(textos), "notas": notas})
    return diapos


def _leer_pptx(ruta):
    """Respaldo sin PowerPoint: lee el archivo directamente."""
    try:
        from pptx import Presentation
    except ImportError:
        return []
    diapos = []
    try:
        prs = Presentation(ruta)
    except Exception:
        return []
    for i, s in enumerate(prs.slides, 1):
        titulo = _limpio(s.shapes.title.text) if s.shapes.title is not None else ""
        textos = []
        for sh in s.shapes:
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                t = _limpio(sh.text_frame.text)
                if t and t != titulo:
                    textos.append(t)
        notas = ""
        if s.has_notes_slide:
            notas = _limpio(s.notes_slide.notes_text_frame.text)
        diapos.append({"numero": i, "titulo": titulo, "texto": " / ".join(textos), "notas": notas})
    return diapos


def contenido(pres=None, ruta=None, forzar=False):
    """Lista de diapositivas de la presentación activa (se guarda para no releerla)."""
    if pres is None and ruta is None:
        pres = _presentacion_activa()
    if pres is not None:
        try:
            ruta = pres.FullName
        except Exception:
            pass
    if not ruta:
        return []
    with _lock:
        if not forzar and ruta in _contenido:
            return _contenido[ruta]
    diapos = []
    if pres is not None:
        try:
            diapos = _leer_com(pres)
        except Exception as e:
            print(f"[No pude leer la presentación por COM: {type(e).__name__}]")
    if not diapos and os.path.isfile(ruta):
        diapos = _leer_pptx(ruta)
    with _lock:
        _contenido[ruta] = diapos
    return diapos


def posicion():
    """(actual, total) si se está exponiendo, si no (None, total)."""
    v = _vista()
    pres = _presentacion_activa()
    total = None
    try:
        total = pres.Slides.Count if pres is not None else None
    except Exception:
        pass
    try:
        return (v.CurrentShowPosition if v is not None else None), total
    except Exception:
        return None, total


def _formatear(d, con_notas=True):
    partes = [f"Diapositiva {d['numero']}"]
    if d["titulo"]:
        partes.append(f"título: {d['titulo']}")
    if d["texto"]:
        partes.append(f"contenido: {d['texto']}")
    if con_notas and d["notas"]:
        partes.append(f"notas del expositor: {d['notas']}")
    return ". ".join(partes)


def contexto(max_chars=6000):
    """Texto para el prompt de sistema: qué presentación hay, en qué diapositiva va y qué dice
    cada una. Así Jarvis puede responder preguntas del público sobre el tema."""
    try:
        pres = _presentacion_activa()
        if pres is None:
            return ""
        nombre = pres.Name
        diapos = contenido(pres)
    except Exception:
        return ""
    actual, total = posicion()
    cab = (f"\n\nPRESENTACIÓN ABIERTA: '{nombre}', {total or len(diapos)} diapositivas"
           + (f", el usuario está exponiendo y va en la {actual}." if actual else " (no está en pantalla completa)."))
    cab += (" Conoces su contenido (abajo). Si te piden explicar una diapositiva, hablar del "
            "tema o responder una pregunta del público sobre la exposición, apóyate en esto. "
            "Para moverte entre diapositivas usa la herramienta presentacion.\n")
    cuerpo = "\n".join(_formatear(d) for d in diapos)
    if len(cuerpo) > max_chars:
        cuerpo = cuerpo[:max_chars] + " [...]"
    return cab + cuerpo


# ---------- Respaldo por teclado ----------
def _teclas(*teclas, escribir=None):
    """Sin COM: se manda la tecla a la presentación (o a lo que esté al frente: sirve para
    Google Slides en el navegador o un PDF a pantalla completa)."""
    import control
    h = _ventana_clase("screenClass", "PPTFrameClass")
    if h:
        control.traer_al_frente(h)
    if escribir:
        pyautogui.write(escribir)
    for t in teclas:
        pyautogui.press(t)


def _buscar_archivo_presentacion(consulta):
    res = archivos._buscar_con_reintento(consulta, "cualquiera", "usuario")
    res = [r for r in res if os.path.splitext(r[3])[1].lower() in EXTENSIONES]
    if not res:
        res = archivos._buscar_con_reintento(consulta, "cualquiera", "equipo")
        res = [r for r in res if os.path.splitext(r[3])[1].lower() in EXTENSIONES]
    return res[0][3] if res else None


# ---------- Skills ----------
@skill("abrir_presentacion",
       "Busca y abre una presentación de PowerPoint por su nombre aproximado ('abre mi "
       "presentación de residencias'). Con iniciar=true además la pone en pantalla completa. "
       "Úsala en vez de abrir_archivo para presentaciones: así Jarvis conoce su contenido.",
       {"consulta": {"type": "string", "description": "Nombre aproximado del archivo"},
        "iniciar": {"type": "boolean", "description": "true para iniciarla en pantalla completa"}},
       requeridos=["consulta"])
def abrir_presentacion(consulta, iniciar=False):
    ruta = _buscar_archivo_presentacion(consulta)
    if not ruta:
        return f"No encontré ninguna presentación parecida a '{consulta}'."
    nombre = os.path.basename(ruta)
    app = _app(crear=True)
    if app is None:  # sin PowerPoint por COM: se abre con lo que tenga Windows asociado
        os.startfile(ruta)
        contenido(ruta=ruta)
        if iniciar:
            time.sleep(4)
            _teclas("f5")
        return f"Abriendo {nombre}."
    pres = None
    try:
        for i in range(1, app.Presentations.Count + 1):
            p = app.Presentations.Item(i)
            if os.path.normcase(p.FullName) == os.path.normcase(ruta):
                pres = p
                break
        if pres is None:
            pres = app.Presentations.Open(ruta, False, False, True)  # ReadOnly, Untitled, WithWindow
        try:
            pres.Windows.Item(1).Activate()
        except Exception:
            pass
    except Exception as e:
        return f"No pude abrir {nombre} en PowerPoint: {type(e).__name__}."
    diapos = contenido(pres, forzar=True)
    if iniciar:
        try:
            pres.SlideShowSettings.Run()
            return f"Presentación {nombre} iniciada, {len(diapos)} diapositivas."
        except Exception:
            _teclas("f5")
    return f"Abrí {nombre}, {len(diapos)} diapositivas."


@skill("presentacion",
       "Controla la presentación de PowerPoint: iniciar (desde el principio), iniciar_aqui "
       "(desde la diapositiva actual), siguiente, anterior, ir (a la diapositiva 'numero'), "
       "primera, ultima, pantalla_negra, pantalla_blanca, reanudar (quita la pantalla negra o "
       "blanca) o terminar.",
       {"accion": {"type": "string",
                   "enum": ["iniciar", "iniciar_aqui", "siguiente", "anterior", "ir", "primera",
                            "ultima", "pantalla_negra", "pantalla_blanca", "reanudar", "terminar"],
                   "description": "Qué hacer"},
        "numero": {"type": "integer", "description": "Para 'ir': número de diapositiva. Para "
                                                     "siguiente/anterior: cuántas (1 por defecto)"}},
       requeridos=["accion"])
def presentacion(accion, numero=0):
    numero = int(numero or 0)
    try:
        return _presentacion_com(accion, numero)
    except _SinCom:
        pass
    except Exception as e:
        print(f"[COM de PowerPoint falló ({type(e).__name__}); uso el teclado]")
    return _presentacion_teclas(accion, numero)


class _SinCom(Exception):
    pass


def _presentacion_com(accion, numero):
    app = _app()
    if app is None:
        raise _SinCom()
    if accion == "iniciar" and _vista() is not None:  # ya estaba en pantalla completa
        _vista().First()
        return Callado("De vuelta a la primera diapositiva.")
    if accion in ("iniciar", "iniciar_aqui"):
        pres = _presentacion_activa()
        if pres is None:
            return "No hay ninguna presentación abierta. Pídeme que abra una."
        desde = 1
        if accion == "iniciar_aqui":
            try:
                desde = app.ActiveWindow.View.Slide.SlideIndex
            except Exception:
                desde = 1
        pres.SlideShowSettings.Run()
        if desde > 1:
            time.sleep(0.4)
            app.SlideShowWindows.Item(1).View.GotoSlide(desde)
        contenido(pres)
        return Callado(f"Presentación iniciada en la diapositiva {desde}.")

    v = _vista()
    if v is None:
        return "La presentación no está en pantalla completa. Dime 'inicia la presentación'."
    total = app.SlideShowWindows.Item(1).Presentation.Slides.Count
    veces = max(1, numero) if accion in ("siguiente", "anterior") else 1

    if accion in ("siguiente", "anterior", "ir", "primera", "ultima") and v.State in (NEGRA, BLANCA):
        v.State = CORRIENDO
    if accion == "siguiente":
        for _ in range(veces):
            v.Next()
    elif accion == "anterior":
        for _ in range(veces):
            v.Previous()
    elif accion == "ir":
        if not 1 <= numero <= total:
            return f"La presentación tiene {total} diapositivas; no existe la {numero}."
        v.GotoSlide(numero)
    elif accion == "primera":
        v.First()
    elif accion == "ultima":
        v.Last()
    elif accion == "pantalla_negra":
        v.State = NEGRA
        return Callado("Pantalla en negro.")
    elif accion == "pantalla_blanca":
        v.State = BLANCA
        return Callado("Pantalla en blanco.")
    elif accion == "reanudar":
        v.State = CORRIENDO
    elif accion == "terminar":
        v.Exit()
        return Callado("Presentación terminada.")
    else:
        return f"No conozco la acción '{accion}'."
    try:
        pos = v.CurrentShowPosition
    except Exception:  # tras pasar la última, PowerPoint muestra "fin de la presentación"
        return Callado("Fin de la presentación.")
    return Callado(f"Diapositiva {pos} de {total}.")


def _presentacion_teclas(accion, numero):
    veces = max(1, numero) if accion in ("siguiente", "anterior") else 1
    mapa = {"iniciar": ("f5",), "iniciar_aqui": ("shift+f5",), "siguiente": ("right",) * veces,
            "anterior": ("left",) * veces, "primera": ("home",), "ultima": ("end",),
            "pantalla_negra": ("b",), "pantalla_blanca": ("w",), "reanudar": ("b",),
            "terminar": ("esc",)}
    if accion == "ir":
        if numero < 1:
            return "Dime a qué número de diapositiva voy."
        _teclas("enter", escribir=str(numero))
        return Callado(f"Diapositiva {numero}.")
    if accion not in mapa:
        return f"No conozco la acción '{accion}'."
    for t in mapa[accion]:
        if "+" in t:
            import control
            h = _ventana_clase("screenClass", "PPTFrameClass")
            if h:
                control.traer_al_frente(h)
            pyautogui.hotkey(*t.split("+"))
        else:
            _teclas(t)
    return Callado("Hecho.")


@skill("explicar_diapositiva",
       "Devuelve el contenido (título, texto y notas del expositor) de la diapositiva actual o "
       "de la indicada, para que puedas explicarla o hablar de ella con el público.",
       {"numero": {"type": "integer", "description": "Número de diapositiva; 0 = la actual"}},
       requeridos=[], terminal=False)
def explicar_diapositiva(numero=0):
    diapos = contenido()
    if not diapos:
        return "No hay ninguna presentación abierta o no pude leer su contenido."
    numero = int(numero or 0)
    if not numero:
        actual, _t = posicion()
        numero = actual or 1
    if not 1 <= numero <= len(diapos):
        return f"La presentación tiene {len(diapos)} diapositivas."
    return _formatear(diapos[numero - 1])
